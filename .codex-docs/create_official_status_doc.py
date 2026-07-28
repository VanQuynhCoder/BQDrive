from __future__ import annotations

import hashlib
import sys
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


EXPECTED_SOURCE_HASH = (
    "51F0AD315F7D90E37FDA727C474A3B1F68A7A87153B282A55D883D58E1EDDE46"
)
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}

REPLACEMENTS = {
    (3, 66, 5): (
        "Trạng thái vòng đời và kiểm duyệt xe: PENDING, APPROVED, RENTED hoặc "
        "REJECTED. Việc ẩn xe được quản lý độc lập bằng isHidden, "
        "hiddenByOwner và hiddenByAdmin."
    ),
    (5, 70, 5): (
        "Trạng thái vòng đời booking: REQUESTED, OWNER_APPROVED, "
        "PAYMENT_PENDING, PAID, IN_PROGRESS, RETURN_INSPECTION, "
        "AWAITING_EXTRA_CHARGE, COMPLETED, CANCELLED, REJECTED hoặc NO_SHOW."
    ),
    (6, 20, 5): (
        "Trạng thái thanh toán tổng hợp: UNPAID, PENDING, DEPOSIT_PAID, "
        "PARTIAL hoặc PAID_FULL."
    ),
    (6, 25, 5): (
        "Trạng thái hợp đồng: ACTIVE, COMPLETED hoặc CANCELLED."
    ),
    (7, 7, 5): (
        "Trạng thái giao dịch thanh toán: PENDING, PAID hoặc FAILED."
    ),
    (7, 12, 5): (
        "Mức độ hoàn tiền của giao dịch: NOT_REFUNDED, PARTIALLY_REFUNDED "
        "hoặc REFUNDED; được suy ra từ amount và refundedAmount."
    ),
    (8, 16, 5): (
        "Trạng thái xử lý hoàn tiền: WAITING_FOR_REFUND_INFO, "
        "MANUAL_REQUIRED, PROCESSING hoặc SUCCEEDED."
    ),
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def replace_cell_text(
    root: etree._Element,
    table_index: int,
    row_index: int,
    cell_index: int,
    value: str,
) -> None:
    tables = root.xpath(".//w:body/w:tbl", namespaces=NS)
    rows = tables[table_index].xpath("./w:tr", namespaces=NS)
    cells = rows[row_index].xpath("./w:tc", namespaces=NS)
    text_nodes = cells[cell_index].xpath(".//w:t", namespaces=NS)
    if not text_nodes:
        raise RuntimeError(
            f"Không tìm thấy text node tại {table_index}/{row_index}/{cell_index}"
        )

    text_nodes[0].text = value
    for node in text_nodes[1:]:
        node.text = ""


def create_document(source: Path, output: Path) -> None:
    with zipfile.ZipFile(source, "r") as source_zip:
        root = etree.fromstring(source_zip.read("word/document.xml"))
        for location, value in REPLACEMENTS.items():
            replace_cell_text(root, *location, value)

        updated_xml = etree.tostring(
            root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        with zipfile.ZipFile(output, "w") as output_zip:
            for item in source_zip.infolist():
                content = (
                    updated_xml
                    if item.filename == "word/document.xml"
                    else source_zip.read(item.filename)
                )
                output_zip.writestr(item, content)


def verify_document(source: Path, output: Path) -> None:
    with (
        zipfile.ZipFile(source, "r") as source_zip,
        zipfile.ZipFile(output, "r") as output_zip,
    ):
        assert source_zip.namelist() == output_zip.namelist()
        changed_parts = [
            name
            for name in source_zip.namelist()
            if source_zip.read(name) != output_zip.read(name)
        ]
        assert changed_parts == ["word/document.xml"], changed_parts

        root = etree.fromstring(output_zip.read("word/document.xml"))
        tables = root.xpath(".//w:body/w:tbl", namespaces=NS)
        for (table_index, row_index, cell_index), expected in REPLACEMENTS.items():
            rows = tables[table_index].xpath("./w:tr", namespaces=NS)
            cells = rows[row_index].xpath("./w:tc", namespaces=NS)
            actual = "".join(
                cells[cell_index].xpath(".//w:t/text()", namespaces=NS)
            )
            assert actual == expected

    document = Document(output)
    assert len(document.sections) == 1
    assert len(document.tables) == 15
    assert len(document.tables[3].rows) == 74
    assert len(document.tables[5].rows) == 93
    assert len(document.tables[6].rows) == 31
    assert len(document.tables[7].rows) == 19
    assert len(document.tables[8].rows) == 53


def main() -> None:
    source = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()
    if source == output:
        raise RuntimeError("Không được ghi đè tài liệu nguồn")
    if sha256(source) != EXPECTED_SOURCE_HASH:
        raise RuntimeError("Tài liệu nguồn đã thay đổi; cần kiểm tra lại")

    create_document(source, output)
    verify_document(source, output)

    if sha256(source) != EXPECTED_SOURCE_HASH:
        raise RuntimeError("Tài liệu nguồn bị thay đổi ngoài ý muốn")

    print(f"[PASS] Created: {output}")
    print(f"[PASS] Replacements: {len(REPLACEMENTS)}")
    print("[PASS] Only word/document.xml changed.")
    print("[PASS] Source unchanged.")


if __name__ == "__main__":
    main()
