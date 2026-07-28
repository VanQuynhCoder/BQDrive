from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

from docx import Document


def main() -> None:
    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    document = Document(source)

    data = {
        "source": str(source),
        "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "sections": [
            {
                "page_width": section.page_width,
                "page_height": section.page_height,
                "top_margin": section.top_margin,
                "right_margin": section.right_margin,
                "bottom_margin": section.bottom_margin,
                "left_margin": section.left_margin,
            }
            for section in document.sections
        ],
        "paragraphs": [
            {
                "index": index,
                "style": paragraph.style.name if paragraph.style else None,
                "text": paragraph.text,
            }
            for index, paragraph in enumerate(document.paragraphs)
            if paragraph.text.strip()
        ],
        "tables": [
            {
                "index": table_index,
                "rows": [
                    {
                        "index": row_index,
                        "cells": [
                            {
                                "index": cell_index,
                                "text": cell.text,
                                "styles": [
                                    paragraph.style.name
                                    if paragraph.style
                                    else None
                                    for paragraph in cell.paragraphs
                                ],
                            }
                            for cell_index, cell in enumerate(row.cells)
                        ],
                    }
                    for row_index, row in enumerate(table.rows)
                ],
            }
            for table_index, table in enumerate(document.tables)
        ],
    }
    output.write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
