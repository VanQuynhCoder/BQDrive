from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def load_builder(path: Path):
    spec = importlib.util.spec_from_file_location("physical_data_builder", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def main() -> None:
    output = Path(sys.argv[1])
    builder_path = Path(__file__).resolve().parent / "build_physical_data_docx.py"
    builder = load_builder(builder_path)
    document = Document(output)

    assert len(document.sections) == 1
    assert len(document.tables) == len(builder.SCHEMAS) + 1

    heading_texts = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style and paragraph.style.name in {"Heading 3", "Heading 4"}
    ]
    assert heading_texts[0] == "3.1.3 Dữ liệu ở mức vật lý"
    for collection, _ in builder.SCHEMAS:
        assert f"Bảng {collection}" in heading_texts
    assert "Các chỉ mục vật lý chính" in heading_texts

    expected_header = ["Thuộc tính", "Kiểu", "K", "U", "M", "Diễn giải"]
    for table, (collection, fields) in zip(document.tables[:-1], builder.SCHEMAS):
        assert len(table.rows) == len(fields) + 1, collection
        assert [cell.text for cell in table.rows[0].cells] == expected_header
        assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
        names = [row.cells[0].text for row in table.rows[1:]]
        assert len(names) == len(set(names)), f"Duplicate fields in {collection}"
        for row in table.rows[1:]:
            assert row.cells[0].text
            assert row.cells[1].text
            assert row.cells[5].text
            assert row._tr.trPr.find(qn("w:trHeight")) is None

    index_table = document.tables[-1]
    assert len(index_table.rows) == len(builder.INDEXES) + 1
    assert [cell.text for cell in index_table.rows[0].cells] == [
        "Collection",
        "Trường chỉ mục",
        "Loại",
        "Mục đích",
    ]
    assert index_table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None

    total_fields = sum(len(fields) for _, fields in builder.SCHEMAS)
    print(
        f"OK: {len(builder.SCHEMAS)} collections, {total_fields} field rows, "
        f"{len(builder.INDEXES)} index rows, {len(document.tables)} tables"
    )


if __name__ == "__main__":
    main()
