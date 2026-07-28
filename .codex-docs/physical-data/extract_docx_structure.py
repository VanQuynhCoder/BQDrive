from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_block_items(document: Document):
    body = document.element.body
    paragraph_by_element = {paragraph._p: paragraph for paragraph in document.paragraphs}
    table_by_element = {table._tbl: table for table in document.tables}

    for child in body.iterchildren():
        if child in paragraph_by_element:
            yield "paragraph", paragraph_by_element[child]
        elif child in table_by_element:
            yield "table", table_by_element[child]


def paragraph_record(paragraph: Paragraph, index: int) -> dict:
    return {
        "kind": "paragraph",
        "index": index,
        "style": paragraph.style.name if paragraph.style else None,
        "text": paragraph.text,
    }


def table_record(table: Table, index: int) -> dict:
    return {
        "kind": "table",
        "index": index,
        "style": table.style.name if table.style else None,
        "rows": [
            [cell.text.replace("\n", " | ") for cell in row.cells]
            for row in table.rows
        ],
    }


def main() -> None:
    source = Path(sys.argv[1])
    destination = Path(sys.argv[2])
    document = Document(source)

    records = []
    paragraph_index = 0
    table_index = 0
    for kind, item in iter_block_items(document):
        if kind == "paragraph":
            records.append(paragraph_record(item, paragraph_index))
            paragraph_index += 1
        else:
            records.append(table_record(item, table_index))
            table_index += 1

    payload = {
        "source": str(source),
        "sections": len(document.sections),
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "blocks": records,
    }
    destination.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
