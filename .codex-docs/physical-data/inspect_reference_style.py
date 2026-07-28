from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def length_points(value):
    return None if value is None else round(value.pt, 2)


def paragraph_info(paragraph):
    fmt = paragraph.paragraph_format
    runs = []
    for run in paragraph.runs:
        color = run.font.color.rgb
        runs.append(
            {
                "text": run.text,
                "font": run.font.name,
                "size_pt": length_points(run.font.size),
                "bold": run.bold,
                "italic": run.italic,
                "color": None if color is None else str(color),
            }
        )
    return {
        "text": paragraph.text,
        "style": paragraph.style.name if paragraph.style else None,
        "alignment": None if paragraph.alignment is None else int(paragraph.alignment),
        "left_indent_pt": length_points(fmt.left_indent),
        "right_indent_pt": length_points(fmt.right_indent),
        "first_line_indent_pt": length_points(fmt.first_line_indent),
        "space_before_pt": length_points(fmt.space_before),
        "space_after_pt": length_points(fmt.space_after),
        "line_spacing": fmt.line_spacing,
        "keep_with_next": fmt.keep_with_next,
        "page_break_before": fmt.page_break_before,
        "runs": runs,
    }


def main():
    document = Document(Path(sys.argv[1]))
    table = document.tables[8]
    table_xml = table._tbl
    grid_cols = [
        int(node.get(qn("w:w")))
        for node in table_xml.tblGrid.findall(qn("w:gridCol"))
    ]
    table_width = table_xml.tblPr.find(qn("w:tblW"))
    table_indent = table_xml.tblPr.find(qn("w:tblInd"))
    cell_margins = table_xml.tblPr.find(qn("w:tblCellMar"))

    payload = {
        "paragraph_532": paragraph_info(document.paragraphs[532]),
        "paragraph_533": paragraph_info(document.paragraphs[533]),
        "normal_style": {
            "font": document.styles["normal"].font.name,
            "size_pt": length_points(document.styles["normal"].font.size),
        },
        "heading_3_style": {
            "font": document.styles["Heading 3"].font.name,
            "size_pt": length_points(document.styles["Heading 3"].font.size),
            "bold": document.styles["Heading 3"].font.bold,
        },
        "heading_4_style": {
            "font": document.styles["Heading 4"].font.name,
            "size_pt": length_points(document.styles["Heading 4"].font.size),
            "bold": document.styles["Heading 4"].font.bold,
        },
        "table_8": {
            "rows": len(table.rows),
            "columns": len(table.columns),
            "grid_dxa": grid_cols,
            "width_type": None
            if table_width is None
            else table_width.get(qn("w:type")),
            "width": None if table_width is None else table_width.get(qn("w:w")),
            "indent_dxa": None
            if table_indent is None
            else float(table_indent.get(qn("w:w"))),
            "cell_margins": None
            if cell_margins is None
            else {
                child.tag.split("}")[-1]: {
                    "type": child.get(qn("w:type")),
                    "width": child.get(qn("w:w")),
                }
                for child in cell_margins
            },
            "header_paragraphs": [
                paragraph_info(cell.paragraphs[0]) for cell in table.rows[0].cells
            ],
            "body_paragraphs": [
                paragraph_info(cell.paragraphs[0]) for cell in table.rows[1].cells
            ],
        },
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
