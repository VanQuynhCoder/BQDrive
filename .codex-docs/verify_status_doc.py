from __future__ import annotations

import hashlib
import sys
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree

from patch_status_doc import EXPECTED_SOURCE_HASH, NS, REPLACEMENTS


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def main() -> None:
    source = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()

    assert sha256(source) == EXPECTED_SOURCE_HASH
    assert source != output

    with (
        zipfile.ZipFile(source, "r") as source_zip,
        zipfile.ZipFile(output, "r") as output_zip,
    ):
        source_names = source_zip.namelist()
        output_names = output_zip.namelist()
        assert source_names == output_names

        changed_parts = []
        for name in source_names:
            if source_zip.read(name) != output_zip.read(name):
                changed_parts.append(name)
        assert changed_parts == ["word/document.xml"], changed_parts

        root = etree.fromstring(output_zip.read("word/document.xml"))
        tables = root.xpath(".//w:body/w:tbl", namespaces=NS)
        for (table_index, row_index, cell_index), expected in REPLACEMENTS.items():
            rows = tables[table_index].xpath("./w:tr", namespaces=NS)
            cells = rows[row_index].xpath("./w:tc", namespaces=NS)
            actual = "".join(
                cells[cell_index].xpath(".//w:t/text()", namespaces=NS)
            )
            assert actual == expected, (
                table_index,
                row_index,
                cell_index,
                actual,
            )

    document = Document(output)
    assert len(document.sections) == 1
    assert len(document.tables) == 15
    assert len(document.tables[3].rows) == 74
    assert len(document.tables[5].rows) == 93
    assert len(document.tables[6].rows) == 31
    assert len(document.tables[7].rows) == 19
    assert len(document.tables[8].rows) == 53

    print("[PASS] DOCX package and status descriptions are valid.")
    print("[PASS] Only word/document.xml changed.")
    print("[PASS] Source hash remains unchanged.")


if __name__ == "__main__":
    main()
