from pathlib import Path
from zipfile import ZipFile

import pdfplumber
from docx import Document


ROOT = Path(__file__).resolve().parent
DOCX = ROOT / "BAOCAO_LVTN_demo.docx"
PDF = ROOT / "Y_MAU_LVTN_2025.pdf"


def extract_docx() -> None:
    doc = Document(DOCX)
    lines = []
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            lines.append(
                f"P{index:04d}\tSTYLE={paragraph.style.name!r}\t{text}"
            )
    for table_index, table in enumerate(doc.tables):
        lines.append(f"\nTABLE {table_index}")
        for row_index, row in enumerate(table.rows):
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            lines.append(f"R{row_index:03d}\t" + " || ".join(cells))
    (ROOT / "docx_text.txt").write_text("\n".join(lines), encoding="utf-8")

    with ZipFile(DOCX) as archive:
        names = set(archive.namelist())
        report = [
            f"paragraphs={len(doc.paragraphs)}",
            f"tables={len(doc.tables)}",
            f"has_comments={'word/comments.xml' in names}",
            f"has_footnotes={'word/footnotes.xml' in names}",
            f"has_endnotes={'word/endnotes.xml' in names}",
        ]
        document_xml = archive.read("word/document.xml")
        report.extend(
            [
                f"tracked_insertions={document_xml.count(b'<w:ins')}",
                f"tracked_deletions={document_xml.count(b'<w:del')}",
                f"explicit_page_breaks={document_xml.count(b'w:type=\"page\"')}",
                f"last_rendered_page_breaks={document_xml.count(b'w:lastRenderedPageBreak')}",
            ]
        )
    (ROOT / "docx_structure.txt").write_text(
        "\n".join(report), encoding="utf-8"
    )
    audit = []
    for style_name in ("Normal", "normal", "Heading 1", "Heading 2", "Heading 3"):
        if style_name not in doc.styles:
            continue
        style = doc.styles[style_name]
        fmt = style.paragraph_format
        audit.append(
            "\t".join(
                [
                    f"STYLE={style_name}",
                    f"font={style.font.name}",
                    f"size_pt={style.font.size.pt if style.font.size else None}",
                    f"bold={style.font.bold}",
                    f"italic={style.font.italic}",
                    f"space_before_pt={fmt.space_before.pt if fmt.space_before else None}",
                    f"space_after_pt={fmt.space_after.pt if fmt.space_after else None}",
                    f"line_spacing={fmt.line_spacing}",
                    f"alignment={fmt.alignment}",
                ]
            )
        )
    def safe_cm(value_getter):
        try:
            value = value_getter()
            return f"{value.cm:.2f}" if value is not None else "None"
        except Exception as exc:
            return f"ERROR({type(exc).__name__}: {exc})"

    for section_index, section in enumerate(doc.sections):
        audit.append(
            "\t".join(
                [
                    f"SECTION={section_index}",
                    f"width_cm={safe_cm(lambda: section.page_width)}",
                    f"height_cm={safe_cm(lambda: section.page_height)}",
                    f"top_cm={safe_cm(lambda: section.top_margin)}",
                    f"bottom_cm={safe_cm(lambda: section.bottom_margin)}",
                    f"left_cm={safe_cm(lambda: section.left_margin)}",
                    f"right_cm={safe_cm(lambda: section.right_margin)}",
                    f"different_first_page={section.different_first_page_header_footer}",
                    f"header_text={' | '.join(p.text for p in section.header.paragraphs)}",
                    f"first_header_text={' | '.join(p.text for p in section.first_page_header.paragraphs)}",
                    f"footer_text={' | '.join(p.text for p in section.footer.paragraphs)}",
                ]
            )
        )
    for paragraph_index in (4, 5, 12, 21, 22, 26, 27):
        paragraph = doc.paragraphs[paragraph_index]
        run_values = []
        for run in paragraph.runs:
            run_values.append(
                {
                    "text": run.text,
                    "font": run.font.name,
                    "size_pt": run.font.size.pt if run.font.size else None,
                    "bold": run.bold,
                    "italic": run.italic,
                    "underline": run.underline,
                }
            )
        audit.append(
            f"PARAGRAPH={paragraph_index}\tstyle={paragraph.style.name}\t"
            f"alignment={paragraph.alignment}\truns={run_values}"
        )
    (ROOT / "docx_format_audit.txt").write_text(
        "\n".join(audit), encoding="utf-8"
    )


def extract_pdf() -> None:
    lines = []
    with pdfplumber.open(PDF) as pdf:
        lines.append(f"TOTAL_PAGES={len(pdf.pages)}")
        for page_number, page in enumerate(pdf.pages, start=1):
            text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
            lines.append(f"\n===== PDF PAGE {page_number} =====\n{text}")
    (ROOT / "pdf_text.txt").write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    extract_docx()
    extract_pdf()
