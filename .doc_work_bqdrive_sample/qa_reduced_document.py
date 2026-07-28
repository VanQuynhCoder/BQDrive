from pathlib import Path
from zipfile import ZipFile
import json
import re

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r"D:\DoAn_LuanVanTotNghiep")
OUTPUT = ROOT / "BQDrive_Nghiep_vu_rut_gon_theo_mau.docx"


def style_name(paragraph):
    return paragraph.style.name if paragraph.style is not None else ""


def has_numbering(paragraph):
    ppr = paragraph._p.pPr
    return ppr is not None and ppr.numPr is not None


def numbering_level(paragraph):
    if not has_numbering(paragraph):
        return None
    return int(paragraph._p.pPr.numPr.ilvl.val)


def main():
    doc = Document(OUTPUT)
    paragraphs = list(doc.paragraphs)
    h1 = [p for p in paragraphs if style_name(p) == "Heading 1"]
    h2 = [p for p in paragraphs if style_name(p) == "Heading 2"]
    labels = [p for p in paragraphs if p.text.strip() == "Các bước thực hiện:"]
    text = "\n".join(p.text for p in paragraphs)

    process_checks = {}
    for index, heading in enumerate(paragraphs):
        if heading not in h2:
            continue
        end = next(
            (
                i for i in range(index + 1, len(paragraphs))
                if style_name(paragraphs[i]) in {"Heading 1", "Heading 2"}
            ),
            len(paragraphs),
        )
        block = paragraphs[index + 1:end]
        main_steps = [p for p in block if numbering_level(p) == 0]
        branches = [p for p in block if numbering_level(p) == 1]
        process_checks[heading.text.strip()] = {
            "lead_paragraphs": len([
                p for p in block
                if p.text.strip()
                and not has_numbering(p)
                and p.text.strip() != "Các bước thực hiện:"
            ]),
            "step_label_count": sum(p.text.strip() == "Các bước thực hiện:" for p in block),
            "main_step_count": len(main_steps),
            "branch_count": len(branches),
            "main_steps_numbered": all(
                re.match(r"^Bước \d+: ", p.text.strip()) for p in main_steps
            ),
        }

    with ZipFile(OUTPUT) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        settings = archive.read("word/settings.xml").decode("utf-8")
        footer_xml = "\n".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if re.fullmatch(r"word/footer\d+\.xml", name)
        )

    section = doc.sections[0]
    normal = doc.styles["Normal"]
    heading1 = doc.styles["Heading 1"]
    heading2 = doc.styles["Heading 2"]
    prohibited = [
        "PRIVATE_OWNER", "CUSTOMER", "đơn hàng", "sản phẩm",
        "Mục đích", "Tác nhân tham gia", "Tiền điều kiện", "Hậu điều kiện",
        "Dữ liệu đầu vào", "Dữ liệu đầu ra", "UseCase", "Usecase",
    ]
    required = [
        "ADMIN", "BUSINESS", "USER", "Google OAuth", "pricing snapshot",
        "NO_SHOW", "30 phút", "WAITING_FOR_REFUND_INFO", "MANUAL_REQUIRED",
        "PROCESSING", "SUCCEEDED", "PAID", "remainingAmount", "IN_PROGRESS",
        "RETURN_INSPECTION", "AWAITING_EXTRA_CHARGE", "COMPLETED",
        "bản đồ", "Task Center", "Notification Center",
    ]
    cancellation_full_markers = [
        p.text for p in paragraphs
        if "Refund bắt đầu ở trạng thái chờ thông tin nhận tiền" in p.text
    ]
    report = {
        "output": str(OUTPUT),
        "file_size": OUTPUT.stat().st_size,
        "group_heading_count": len(h1),
        "group_headings": [p.text.strip() for p in h1],
        "process_heading_count": len(h2),
        "unique_process_heading_count": len({p.text.strip() for p in h2}),
        "step_label_count": len(labels),
        "process_checks": process_checks,
        "processes_outside_step_range": {
            title: data["main_step_count"]
            for title, data in process_checks.items()
            if data["main_step_count"] < 5 or data["main_step_count"] > 14
        },
        "processes_without_branches": [
            title for title, data in process_checks.items() if data["branch_count"] == 0
        ],
        "processes_without_single_label": [
            title for title, data in process_checks.items() if data["step_label_count"] != 1
        ],
        "processes_with_bad_step_numbers": [
            title for title, data in process_checks.items() if not data["main_steps_numbered"]
        ],
        "media_count": len(media),
        "table_count": len(doc.tables),
        "toc_field_count": doc._element.xml.count("TOC "),
        "page_field_count": footer_xml.count("PAGE"),
        "update_fields_on_open": '<w:updateFields w:val="true"' in settings,
        "required_terms_missing": [term for term in required if term not in text],
        "prohibited_terms_found": [term for term in prohibited if term in text],
        "full_refund_flow_marker_count": len(cancellation_full_markers),
        "page": {
            "width_cm": round(section.page_width.cm, 2),
            "height_cm": round(section.page_height.cm, 2),
            "left_cm": round(section.left_margin.cm, 2),
            "right_cm": round(section.right_margin.cm, 2),
            "top_cm": round(section.top_margin.cm, 2),
            "bottom_cm": round(section.bottom_margin.cm, 2),
        },
        "styles": {
            "normal_font": normal.font.name,
            "normal_size_pt": normal.font.size.pt if normal.font.size else None,
            "normal_spacing": normal.paragraph_format.line_spacing,
            "heading1_font": heading1.font.name,
            "heading1_size_pt": heading1.font.size.pt if heading1.font.size else None,
            "heading2_font": heading2.font.name,
            "heading2_size_pt": heading2.font.size.pt if heading2.font.size else None,
        },
    }
    path = ROOT / ".doc_work_bqdrive_sample" / "reduced_qa_report.json"
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
