from pathlib import Path
import json

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"D:\DoAn_LuanVanTotNghiep")
WORK = ROOT / ".doc_work_bqdrive_sample"
REFERENCE = Path(r"C:\Users\bom13\Downloads\VuHoangUng-1.docx")
OUTPUT = ROOT / "BQDrive_Nghiep_vu_rut_gon_theo_mau.docx"
FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)


def proc(title, lead, steps):
    return {"title": title, "lead": lead, "steps": steps}


def main(text):
    return (0, text)


def branch(text):
    return (1, text)


GROUPS = [
    {
        "title": "Nghiệp vụ tài khoản và phân quyền",
        "intro": [
            "Nhóm nghiệp vụ tài khoản và phân quyền quản lý việc hình thành, xác thực và sử dụng danh tính trên BQDrive. Hệ thống sử dụng ba vai trò ADMIN, BUSINESS và USER; trong đó USER có thể tham gia với tư cách người thuê hoặc chủ xe có xe ký gửi, còn BUSINESS đại diện cho doanh nghiệp quản lý đội xe.",
            "USER được tự đăng ký và phải đáp ứng yêu cầu xác thực trước khi sử dụng đầy đủ chức năng. Tài khoản BUSINESS do ADMIN tạo hoặc cấp sau khi kiểm tra hồ sơ doanh nghiệp; mọi tài khoản chỉ được truy cập dữ liệu và thao tác phù hợp với vai trò, trạng thái hoạt động và phạm vi sở hữu.",
        ],
        "processes": [
            proc(
                "Quy trình đăng ký và xác thực tài khoản USER",
                "Quy trình này kết hợp đăng ký, gửi mã OTP và xác thực email trong một luồng thống nhất. Tài khoản chỉ được kích hoạt sau khi email và dữ liệu đăng ký hợp lệ, đồng thời mã OTP được xác nhận trong thời hạn cho phép.",
                [
                    main("Người đăng ký cung cấp họ tên, email, mật khẩu và các thông tin bắt buộc của tài khoản USER."),
                    main("Hệ thống kiểm tra định dạng email, độ đầy đủ của dữ liệu và việc email đã được sử dụng hay chưa."),
                    branch("Nếu email trùng hoặc dữ liệu không hợp lệ, hệ thống từ chối tạo tài khoản và nêu nội dung cần điều chỉnh."),
                    branch("Nếu dữ liệu hợp lệ, hệ thống tạo tài khoản USER ở trạng thái chờ xác thực email."),
                    main("Hệ thống tạo mã OTP có thời hạn và gửi mã đến email đã đăng ký."),
                    main("Người đăng ký nhập OTP để xác nhận quyền sử dụng email."),
                    main("Hệ thống kiểm tra OTP có thuộc đúng tài khoản, còn hiệu lực và chưa vượt quá số lần thử hay không."),
                    branch("Nếu OTP sai hoặc hết hạn, tài khoản vẫn chưa được kích hoạt; người đăng ký có thể yêu cầu gửi lại mã theo giới hạn của hệ thống."),
                    branch("Nếu OTP hợp lệ, hệ thống đánh dấu email đã xác thực và kích hoạt tài khoản."),
                    main("Hệ thống thông báo kết quả xác thực và cho phép USER đăng nhập bằng thông tin đã đăng ký."),
                ],
            ),
            proc(
                "Quy trình đăng nhập và khôi phục tài khoản",
                "Quy trình cho phép người dùng đăng nhập bằng email và mật khẩu hoặc Google OAuth, đồng thời hỗ trợ khôi phục khi quên mật khẩu. Mọi phương thức đều phải kiểm tra trạng thái tài khoản để ngăn tài khoản bị khóa hoặc xóa mềm tiếp tục truy cập.",
                [
                    main("Người dùng chọn đăng nhập bằng email và mật khẩu hoặc chọn Google OAuth."),
                    main("Với email và mật khẩu, hệ thống đối chiếu thông tin xác thực và trạng thái xác thực email."),
                    main("Với Google OAuth, hệ thống xác minh thông tin Google; nếu email chưa tồn tại thì có thể tạo tài khoản USER từ email đã được Google xác thực."),
                    main("Hệ thống kiểm tra tài khoản không bị khóa, không bị xóa mềm và đáp ứng điều kiện xác thực."),
                    branch("Nếu thông tin sai hoặc tài khoản không đủ điều kiện, hệ thống từ chối đăng nhập và không tạo phiên sử dụng."),
                    branch("Nếu hợp lệ, hệ thống tạo phiên theo đúng vai trò ADMIN, BUSINESS hoặc USER."),
                    main("Khi quên mật khẩu, người dùng yêu cầu khôi phục bằng email của tài khoản đang hoạt động."),
                    main("Hệ thống gửi thông tin khôi phục; người dùng xác nhận mã hợp lệ và thiết lập mật khẩu mới."),
                    branch("Nếu mã khôi phục sai hoặc hết hạn, mật khẩu không thay đổi và người dùng phải yêu cầu lại."),
                    main("Sau khi đặt lại thành công, thông tin khôi phục cũ bị vô hiệu hóa và người dùng đăng nhập lại bằng mật khẩu mới."),
                ],
            ),
            proc(
                "Quy trình ADMIN quản lý tài khoản và doanh nghiệp",
                "ADMIN quản lý vòng đời tài khoản và việc cấp quyền cho doanh nghiệp từ một quy trình chung. Các thao tác tìm kiếm, tạo BUSINESS, khóa, mở khóa và xóa mềm đều phải tuân theo quyền quản trị và không làm phát sinh vai trò ngoài ADMIN, BUSINESS và USER.",
                [
                    main("ADMIN xem hoặc tìm kiếm tài khoản theo thông tin nhận diện, vai trò và trạng thái."),
                    main("Khi cấp tài khoản BUSINESS, ADMIN nhập thông tin doanh nghiệp và người đại diện cần sử dụng tài khoản."),
                    main("Hệ thống kiểm tra email không trùng, hồ sơ doanh nghiệp hợp lệ và quyền được cấp phù hợp chính sách."),
                    branch("Nếu dữ liệu hoặc quyền không hợp lệ, hệ thống từ chối tạo BUSINESS và nêu nguyên nhân."),
                    main("Hệ thống tạo hoặc liên kết tài khoản với hồ sơ doanh nghiệp, sau đó gán vai trò BUSINESS."),
                    main("Khi cần kiểm soát truy cập, ADMIN chọn khóa, mở khóa hoặc xóa mềm tài khoản phù hợp."),
                    main("Hệ thống kiểm tra ADMIN có quyền thao tác và không cho phép cấp quyền trái chính sách hệ thống."),
                    branch("Tài khoản bị khóa hoặc xóa mềm không được đăng nhập; mở khóa chỉ khôi phục truy cập khi các điều kiện khác vẫn hợp lệ."),
                    main("Hệ thống lưu trạng thái mới và thông báo thay đổi cho tài khoản chịu tác động."),
                ],
            ),
        ],
    },
    {
        "title": "Nghiệp vụ quản lý và kiểm duyệt xe",
        "intro": [
            "Nhóm nghiệp vụ này quản lý toàn bộ hồ sơ xe do BUSINESS hoặc USER chủ xe ký gửi sở hữu. Hồ sơ bao gồm thông tin nhận diện, thương hiệu, biển số, thông số kỹ thuật, hình ảnh, cấu hình giá và trạng thái hiển thị; chủ xe chỉ được thao tác trên xe thuộc phạm vi quản lý của mình.",
            "Xe mới hoặc xe thay đổi thông tin quan trọng phải được ADMIN kiểm duyệt trước khi công khai. Vị trí nhận xe được quản lý riêng bằng địa chỉ và tọa độ bản đồ; việc chỉ cập nhật vị trí có thể giữ trạng thái đã duyệt của xe theo nghiệp vụ hiện tại.",
        ],
        "processes": [
            proc(
                "Quy trình chủ xe quản lý thông tin xe",
                "Quy trình gộp các thao tác tạo, cập nhật, quản lý hình ảnh, cấu hình giá, ẩn, hiện và xóa mềm xe. Mọi thay đổi đều gắn với quyền sở hữu và phải bảo vệ các booking, hợp đồng hoặc lượt giữ đang còn hiệu lực.",
                [
                    main("Chủ xe mở phạm vi quản lý xe của BUSINESS hoặc USER chủ xe ký gửi tương ứng."),
                    main("Khi tạo xe, chủ xe nhập biển số, thương hiệu, thông số kỹ thuật, hình ảnh, chính sách thuê và cấu hình giá."),
                    main("Hệ thống kiểm tra quyền sở hữu, biển số, dữ liệu bắt buộc, hình ảnh và các mức giá hợp lệ."),
                    branch("Nếu dữ liệu không hợp lệ hoặc biển số bị trùng, hệ thống từ chối lưu và yêu cầu điều chỉnh."),
                    main("Hệ thống tạo hồ sơ xe và chuyển xe sang trạng thái chờ duyệt (PENDING)."),
                    main("Khi cập nhật, chủ xe chỉ được sửa xe thuộc quyền quản lý; thay đổi quan trọng làm xe quay lại PENDING để kiểm duyệt."),
                    main("Chủ xe có thể thêm, thay đổi hình ảnh và điều chỉnh cấu hình giá trong cùng hồ sơ xe."),
                    main("Khi ẩn, hiển thị hoặc xóa mềm, hệ thống kiểm tra xe không có booking, hợp đồng hoặc lượt giữ đang hoạt động."),
                    branch("Nếu còn công việc thuê có hiệu lực, hệ thống từ chối ẩn, hiện hoặc xóa để không phá vỡ giao dịch đang xử lý."),
                    main("Nếu đủ điều kiện, hệ thống cập nhật trạng thái hiển thị hoặc cờ xóa mềm và loại xe khỏi phạm vi công khai khi cần."),
                ],
            ),
            proc(
                "Quy trình quản lý vị trí nhận xe",
                "Vị trí nhận xe hỗ trợ người thuê xem địa điểm và lập kế hoạch nhận xe trên bản đồ. Hệ thống lưu đồng thời địa chỉ mô tả, vĩ độ và kinh độ để bảo đảm thông tin vị trí có thể tái sử dụng trong tìm kiếm, booking và hợp đồng.",
                [
                    main("Chủ xe chọn xe thuộc quyền quản lý và nhập hoặc tìm kiếm địa chỉ nhận xe."),
                    main("Hệ thống hiển thị kết quả địa chỉ để chủ xe chọn điểm phù hợp trên bản đồ."),
                    main("Chủ xe điều chỉnh ghim bản đồ khi vị trí tự động chưa chính xác."),
                    main("Hệ thống kiểm tra vĩ độ nằm trong khoảng từ -90 đến 90, kinh độ nằm trong khoảng từ -180 đến 180 và quyền quản lý xe hợp lệ."),
                    branch("Nếu địa chỉ hoặc tọa độ không hợp lệ, hệ thống không lưu và yêu cầu chọn lại vị trí."),
                    main("Nếu hợp lệ, hệ thống lưu địa chỉ, tọa độ, ghi chú và lịch sử vị trí nhận xe."),
                    main("Hệ thống dùng vị trí đã lưu để hiển thị bản đồ và tạo bản chụp địa điểm khi phát sinh booking."),
                    main("Nếu xe đã được duyệt và chỉ thay đổi vị trí nhận xe, xe có thể giữ trạng thái đã duyệt (APPROVED)."),
                ],
            ),
            proc(
                "Quy trình ADMIN kiểm duyệt xe",
                "ADMIN kiểm soát chất lượng hồ sơ trước khi xe được công khai. Quy trình bao gồm duyệt, từ chối, lưu lý do và xử lý yêu cầu duyệt lại sau khi chủ xe đã chỉnh sửa nội dung bị từ chối.",
                [
                    main("ADMIN xem danh sách xe đang chờ duyệt (PENDING) và mở hồ sơ cần kiểm tra."),
                    main("Hệ thống cung cấp thông tin xe, thương hiệu, hình ảnh, biển số, cấu hình giá, vị trí và chủ sở hữu."),
                    main("ADMIN kiểm tra tính đầy đủ, hợp lệ và sự phù hợp của hồ sơ với chính sách hệ thống."),
                    branch("Nếu hồ sơ đạt yêu cầu, ADMIN chọn duyệt xe."),
                    branch("Nếu hồ sơ chưa đạt, ADMIN chọn từ chối và bắt buộc nhập lý do cụ thể."),
                    main("Xe được duyệt chuyển sang APPROVED; xe bị từ chối chuyển sang trạng thái bị từ chối (REJECTED) và lưu lý do."),
                    main("Chủ xe chỉnh sửa nội dung theo lý do từ chối và gửi lại để xe quay về PENDING."),
                    main("Hệ thống chỉ công khai xe APPROVED, không bị ẩn và không bị xóa mềm; xe chưa duyệt không xuất hiện trong kết quả thuê công khai."),
                    main("Kết quả kiểm duyệt được gửi đến chủ xe để tiếp tục quản lý hồ sơ."),
                ],
            ),
        ],
    },
    {
        "title": "Nghiệp vụ tìm kiếm, giỏ xe và vòng đời booking",
        "intro": [
            "Nhóm nghiệp vụ này bao quát giai đoạn từ lúc người thuê tìm xe đến khi booking được chủ xe quyết định có tiếp tục hay không. Người thuê có thể giữ xe tạm thời trong giỏ hoặc tạo booking trực tiếp; cả hai cách đều phải kiểm tra quyền đặt, thời gian thuê và khả năng đáp ứng của xe.",
            "Booking mới gắn với đúng người thuê, xe và chủ xe là BUSINESS hoặc USER chủ xe ký gửi. Nhóm này chỉ mô tả từ chối và NO_SHOW ở mức vòng đời; luồng chủ động hủy booking và hoàn tiền được trình bày duy nhất tại quy trình 2.4.1.4.3.",
        ],
        "processes": [
            proc(
                "Quy trình tìm kiếm và xem chi tiết xe",
                "Quy trình hỗ trợ người thuê xác định xe phù hợp theo địa điểm, thời gian và các tiêu chí hiện có. Giá hiển thị tại đây là giá dự kiến theo chính sách giá; công thức tính và lưu giá được trình bày tại quy trình 2.4.1.4.1.",
                [
                    main("Người thuê nhập địa điểm, thời gian bắt đầu, thời gian kết thúc và các tiêu chí tìm kiếm."),
                    main("Hệ thống kiểm tra thời gian bắt đầu phải trước thời gian kết thúc và dữ liệu tìm kiếm hợp lệ."),
                    branch("Nếu khoảng thời gian không hợp lệ, hệ thống từ chối tìm kiếm theo thời gian và yêu cầu điều chỉnh."),
                    main("Hệ thống lọc theo thương hiệu, mức giá, số ghế, nhiên liệu, hộp số và các tiêu chí đang được hỗ trợ."),
                    main("Hệ thống chỉ lấy xe đã APPROVED, không bị ẩn, không bị xóa mềm và không thuộc chính người thuê."),
                    main("Hệ thống loại các xe trùng lịch hoặc không khả dụng trong khoảng thời gian đã chọn."),
                    branch("Nếu không có kết quả, hệ thống thông báo và cho phép thay đổi địa điểm, thời gian hoặc bộ lọc."),
                    main("Người thuê mở chi tiết để xem hình ảnh, đặc điểm xe, giá dự kiến, đánh giá và vị trí nhận xe trên bản đồ."),
                    main("Người thuê chọn xe phù hợp để giữ tạm thời hoặc tạo booking trực tiếp."),
                ],
            ),
            proc(
                "Quy trình thêm xe vào giỏ và giữ xe tạm thời",
                "Giỏ xe tạo một khoảng giữ có thời hạn để người thuê hoàn thiện quyết định trước khi tạo booking. Lượt giữ có bản chụp giá riêng và không được dùng sau khi hết hạn hoặc khi quyền sở hữu không khớp.",
                [
                    main("Người thuê chọn xe và khoảng thời gian cần thuê để thêm vào giỏ."),
                    main("Hệ thống kiểm tra xe còn hợp lệ, người thuê không phải chủ xe và khoảng thời gian không trùng lịch."),
                    main("Hệ thống kiểm tra không có lượt giữ đang hoạt động khác xung đột theo điều kiện khóa lịch hiện tại."),
                    branch("Nếu xe hoặc thời gian không khả dụng, hệ thống không tạo mục giỏ và nêu nguyên nhân."),
                    main("Hệ thống tính giá dự kiến theo quy trình giá và lưu pricing snapshot tại thời điểm giữ xe."),
                    main("Hệ thống tạo mục giỏ ở trạng thái đang hoạt động (ACTIVE) và thiết lập thời điểm hết hạn."),
                    main("Trong thời hạn giữ, khoảng thời gian được bảo vệ để hạn chế phát sinh lượt giữ hoặc booking xung đột."),
                    branch("Nếu người thuê tiếp tục đúng hạn, hệ thống cho phép chuyển mục giỏ thành booking."),
                    branch("Nếu hết hạn, hệ thống chuyển lượt giữ sang hết hiệu lực, giải phóng khoảng thời gian và không cho tạo booking từ mục giỏ đó."),
                    main("Người thuê có thể chủ động hủy mục giỏ chưa sử dụng để giải phóng lượt giữ sớm."),
                ],
            ),
            proc(
                "Quy trình tạo booking trực tiếp hoặc từ giỏ",
                "Booking có thể được tạo trực tiếp từ chi tiết xe hoặc từ một mục giỏ còn hiệu lực. Hai nguồn sử dụng cùng các quy tắc cốt lõi về người thuê, xe, thời gian, chủ sở hữu và bản chụp giá.",
                [
                    main("Người thuê chọn tạo booking trực tiếp hoặc tiếp tục từ mục giỏ."),
                    main("Hệ thống kiểm tra tài khoản người thuê hoạt động và thông tin người thuê đã được cung cấp đầy đủ."),
                    main("Hệ thống kiểm tra xe tồn tại, APPROVED, không bị ẩn, không bị xóa mềm và không thuộc chính người thuê."),
                    main("Hệ thống kiểm tra thời gian bắt đầu trước thời gian kết thúc và không trùng booking hoặc lượt giữ có hiệu lực."),
                    main("Nếu tạo từ giỏ, hệ thống kiểm tra mục giỏ thuộc đúng người thuê, còn ACTIVE và chưa hết hạn."),
                    branch("Nếu bất kỳ điều kiện nào không đạt, hệ thống từ chối tạo booking và giữ nguyên dữ liệu hiện có."),
                    main("Hệ thống xác định chủ xe là BUSINESS hoặc USER chủ xe ký gửi theo hồ sơ sở hữu xe."),
                    main("Hệ thống lưu người thuê, xe, chủ xe, thời gian, thông tin nhận xe và pricing snapshot."),
                    main("Hệ thống tạo booking ở trạng thái chờ chủ xe duyệt (REQUESTED)."),
                    main("Nếu booking được tạo từ giỏ, hệ thống đánh dấu mục giỏ đã được sử dụng và không cho tái sử dụng."),
                    main("Hệ thống tạo việc cần làm và gửi thông báo cho đúng chủ xe."),
                ],
            ),
            proc(
                "Quy trình chủ xe duyệt hoặc từ chối booking",
                "Chủ xe quyết định yêu cầu thuê có thể tiếp tục dựa trên quyền quản lý xe, thời gian và khả năng bàn giao. Quy trình kết thúc bằng một trong hai nhánh rõ ràng: duyệt để chuyển sang hợp đồng và thanh toán, hoặc từ chối kèm lý do.",
                [
                    main("Chủ xe mở booking đang chờ xử lý liên quan đến xe thuộc quyền quản lý."),
                    main("Hệ thống kiểm tra booking đang REQUESTED và chủ xe đúng là BUSINESS hoặc USER chủ xe ký gửi của xe."),
                    branch("Nếu không đúng quyền hoặc trạng thái, hệ thống từ chối thao tác."),
                    main("Chủ xe kiểm tra thời gian thuê, thông tin người thuê, lịch xe và khả năng bàn giao."),
                    branch("Nếu chấp nhận, chủ xe duyệt booking; booking chuyển sang trạng thái chủ xe đã duyệt (OWNER_APPROVED)."),
                    branch("Nếu không chấp nhận, chủ xe nhập lý do; booking chuyển sang REJECTED và không tiếp tục thực hiện."),
                    main("Hệ thống cập nhật Task Center và gửi kết quả đến người thuê."),
                    main("Booking OWNER_APPROVED được chuyển sang giai đoạn tạo hoặc hoàn thiện hợp đồng và thực hiện nghĩa vụ thanh toán."),
                    main("Nếu việc từ chối phát sinh vấn đề tài chính, hệ thống dẫn chiếu xử lý theo quy trình hủy booking và hoàn tiền 2.4.1.4.3."),
                ],
            ),
            proc(
                "Quy trình xử lý booking không được thực hiện và trường hợp NO_SHOW",
                "Quy trình xử lý booking không tiếp tục do bị từ chối, mất điều kiện thực hiện hoặc người thuê không đến nhận xe. Nội dung không lặp lại luồng người thuê hoặc chủ xe chủ động hủy booking; vấn đề hoàn tiền, nếu có, được chuyển sang nhóm tài chính.",
                [
                    main("Hệ thống xác định booking không thể tiếp tục do bị từ chối hoặc không còn đáp ứng điều kiện thực hiện."),
                    main("Đối với trường hợp người thuê không đến nhận xe, chủ xe mở booking thuộc quyền quản lý và yêu cầu ghi nhận NO_SHOW."),
                    main("Hệ thống kiểm tra booking thuộc trạng thái cho phép, xe chưa được bàn giao và thời điểm hiện tại đã qua giờ nhận xe."),
                    main("Hệ thống tính thời điểm được phép ghi nhận bằng giờ nhận xe cộng khoảng chờ 30 phút."),
                    branch("Nếu chưa đủ 30 phút, hệ thống từ chối thao tác NO_SHOW và booking giữ nguyên."),
                    branch("Nếu đã đủ thời gian và các điều kiện còn hợp lệ, chủ xe nhập lý do người thuê không đến."),
                    main("Hệ thống chuyển booking sang trạng thái khách không đến nhận xe (NO_SHOW) và ghi nhận thời điểm xử lý."),
                    main("Hệ thống giải phóng lịch hoặc cập nhật khả năng sử dụng xe theo trạng thái nghiệp vụ."),
                    main("Hệ thống cập nhật việc cần làm và thông báo cho các bên liên quan."),
                    main("Nếu booking có payment cần xử lý, nghiệp vụ tài chính được chuyển sang quy trình 2.4.1.4.3 mà không mô tả lại tại đây."),
                ],
            ),
        ],
    },
    {
        "title": "Nghiệp vụ tính giá, thanh toán và hoàn tiền",
        "intro": [
            "Nhóm nghiệp vụ tài chính xác định số tiền thuê, ghi nhận các khoản thanh toán và xử lý hoàn tiền khi booking bị hủy hợp lệ. Giá được tính theo chế độ thuê giờ hoặc ngày, theo từng khoảng thời gian và ưu tiên ngày lễ trước cuối tuần, cuối tuần trước ngày thường.",
            "Payment và Refund có vòng đời độc lập với booking. Chỉ payment đã thanh toán thành công mới được cộng vào số tiền đã thanh toán; hoàn tiền hiện được thực hiện thủ công theo hồ sơ Refund, có thông tin người nhận, xác nhận gửi tiền của chủ xe và xác nhận đã nhận của người thuê.",
        ],
        "processes": [
            proc(
                "Quy trình tính và lưu giá thuê",
                "Quy trình tạo kết quả giá nhất quán cho giỏ và booking mà không phụ thuộc vào thay đổi giá về sau. Hệ thống sử dụng cấu hình giá của xe và dữ liệu ngày lễ hiện hành, nhưng không tự áp dụng công thức ngoài các quy tắc đã triển khai.",
                [
                    main("Hệ thống nhận xe, thời gian bắt đầu, thời gian kết thúc và chế độ thuê theo giờ hoặc theo ngày."),
                    main("Hệ thống kiểm tra chế độ thuê được xe cho phép và khoảng thời gian có giá trị."),
                    branch("Nếu dữ liệu hoặc chế độ thuê không hợp lệ, hệ thống không tính giá và yêu cầu điều chỉnh."),
                    main("Hệ thống chia thời gian thuê thành các khoảng cần tính theo logic hiện tại."),
                    main("Với mỗi khoảng, hệ thống xác định có thuộc ngày lễ, cuối tuần hay ngày thường."),
                    main("Hệ thống ưu tiên giá ngày lễ; nếu không phải ngày lễ thì ưu tiên giá cuối tuần; còn lại dùng giá ngày thường."),
                    main("Hệ thống áp dụng đơn giá theo giờ hoặc theo ngày tương ứng và cộng thành tổng giá thuê."),
                    main("Hệ thống tạo pricing snapshot gồm đơn giá, loại ngày, các thành phần tính giá và tổng tiền."),
                    main("Snapshot được lưu cùng lượt giữ hoặc booking để không thay đổi khi chủ xe cập nhật giá sau đó."),
                ],
            ),
            proc(
                "Quy trình thanh toán booking",
                "Quy trình hợp nhất thanh toán cọc, toàn bộ và phần còn lại bằng CASH, MOMO hoặc VNPAY. Payment được theo dõi riêng, chống ghi nhận lặp và chỉ ảnh hưởng số tiền đã trả khi đạt trạng thái thành công.",
                [
                    main("Người thuê mở booking đã được chủ xe duyệt và chọn thanh toán cọc, toàn bộ hoặc phần còn lại."),
                    main("Người thuê chọn phương thức CASH, MOMO hoặc VNPAY theo lựa chọn đang được hỗ trợ."),
                    main("Hệ thống kiểm tra booking còn cho phép thanh toán, số tiền hợp lệ và khoản tương ứng chưa được ghi nhận trùng."),
                    branch("Nếu điều kiện không hợp lệ, hệ thống không tạo giao dịch mới và nêu lý do."),
                    main("Hệ thống tạo payment ở trạng thái chờ xử lý (PENDING) và liên kết với booking."),
                    main("Với MOMO hoặc VNPAY, người thuê thực hiện giao dịch và hệ thống chờ kết quả callback; với CASH, chủ xe xác nhận sau khi thực tế đã thu tiền."),
                    main("Hệ thống kiểm tra mã giao dịch, số tiền, booking và dữ liệu xác thực trước khi cập nhật payment."),
                    branch("Nếu thất bại, payment chuyển sang FAILED và không được cộng vào số tiền đã thanh toán."),
                    branch("Nếu thành công, payment chuyển sang đã thanh toán (PAID) và chỉ được ghi nhận một lần."),
                    main("Hệ thống tính paidAmount từ các payment PAID và tính remainingAmount bằng phần nghĩa vụ còn lại."),
                    main("Thanh toán cọc không đồng nghĩa thanh toán đủ; chỉ khi remainingAmount bằng 0 mới được xem là hoàn tất nghĩa vụ thanh toán."),
                    main("Nếu callback đến sau khi booking đã CANCELLED, hệ thống ghi nhận payment và tạo xử lý hoàn phù hợp nhưng không kích hoạt lại booking."),
                    main("Kết quả payment được phản ánh vào booking, Task Center và thông báo liên quan."),
                ],
            ),
            proc(
                "Quy trình hủy booking và hoàn tiền",
                "Đây là vị trí duy nhất mô tả đầy đủ việc người thuê hoặc chủ xe chủ động hủy booking và xử lý số tiền cần hoàn. Booking chuyển sang đã hủy, còn Refund tiếp tục vòng đời thủ công riêng cho đến khi người thuê xác nhận đã nhận tiền.",
                [
                    main("Người thuê hoặc chủ xe có quyền mở booking, xem trước chính sách và nhập lý do hủy."),
                    main("Hệ thống kiểm tra quyền của người yêu cầu và booking còn ở REQUESTED, OWNER_APPROVED, PAYMENT_PENDING, PAID hoặc trạng thái cũ tương ứng."),
                    branch("Booking đã IN_PROGRESS, RETURN_INSPECTION, AWAITING_EXTRA_CHARGE, COMPLETED, REJECTED hoặc NO_SHOW không được hủy theo luồng thông thường."),
                    main("Hệ thống tính trước phí hủy và số tiền dự kiến hoàn theo thời điểm, bên hủy và bản chụp chính sách của booking."),
                    main("Người yêu cầu xác nhận hủy sau khi xem thông tin tài chính."),
                    main("Hệ thống chuyển booking sang trạng thái đã hủy (CANCELLED), lưu lý do và giải phóng lịch xe."),
                    main("Hệ thống truy xuất các payment PAID hoặc đã hoàn một phần và loại trừ số tiền đã hoàn trước đó."),
                    main("Hệ thống áp dụng chính sách hủy, đồng thời giới hạn tổng tiền hoàn không vượt quá số tiền đã thanh toán hợp lệ còn lại."),
                    branch("Nếu refundAmount bằng 0, hệ thống không tạo Refund và kết thúc phần hoàn tiền."),
                    branch("Nếu refundAmount lớn hơn 0, hệ thống tạo một hồ sơ Refund chống trùng và liên kết booking cùng các payment liên quan."),
                    main("Refund bắt đầu ở trạng thái chờ thông tin nhận tiền (WAITING_FOR_REFUND_INFO); người thuê cung cấp tài khoản ngân hàng, ví điện tử hoặc thông tin nhận tiền mặt hợp lệ."),
                    main("Sau khi có thông tin người nhận, Refund chuyển sang yêu cầu xử lý thủ công (MANUAL_REQUIRED)."),
                    main("Chủ xe chuyển tiền ngoài hệ thống, ghi phương thức, mã tham chiếu, ghi chú hoặc minh chứng; Refund chuyển sang đang xử lý (PROCESSING)."),
                    main("Người thuê kiểm tra và xác nhận đã nhận tiền; khi đó hệ thống cập nhật payment liên quan và chuyển Refund sang hoàn tất (SUCCEEDED)."),
                    main("Hệ thống không cho xác nhận hoặc hoàn trùng; callback payment đến muộn vẫn không làm booking CANCELLED hoạt động trở lại."),
                ],
            ),
        ],
    },
    {
        "title": "Nghiệp vụ bàn giao, trả xe và hoàn tất chuyến thuê",
        "intro": [
            "Nhóm nghiệp vụ này quản lý giai đoạn thực hiện chuyến thuê từ trước lúc bàn giao đến khi xe được trả và mọi nghĩa vụ phát sinh đã được xử lý. Chỉ chủ xe quản lý xe trong booking mới có quyền bàn giao, nhận lại xe và xác nhận hoàn tất.",
            "Hợp đồng, danh tính người thuê và nghĩa vụ thanh toán phải đáp ứng điều kiện trước khi xe được giao. Sau khi xe trả, biên bản kiểm tra là căn cứ xác định phụ phí; booking chỉ hoàn tất khi không còn kiểm tra hoặc khoản thanh toán bắt buộc đang chờ.",
        ],
        "processes": [
            proc(
                "Quy trình kiểm tra và bàn giao xe",
                "Quy trình kết hợp kiểm tra booking, người thuê, hợp đồng, thanh toán và việc bàn giao thực tế. Booking chỉ chuyển sang giai đoạn thuê sau khi chủ xe xác nhận đầy đủ các điều kiện.",
                [
                    main("Chủ xe mở booking đến thời điểm nhận xe và thuộc xe mình quản lý."),
                    main("Hệ thống kiểm tra chủ xe đúng quyền và booking đã được duyệt."),
                    main("Chủ xe đối chiếu người nhận xe với thông tin người thuê đã lưu trong booking và hợp đồng."),
                    main("Hệ thống kiểm tra hợp đồng đã được tạo hoặc hoàn thiện theo yêu cầu."),
                    main("Hệ thống kiểm tra payment PAID, paidAmount và remainingAmount theo điều kiện bàn giao."),
                    branch("Nếu còn khoản bắt buộc, chủ xe yêu cầu người thuê thanh toán hoặc xác nhận tiền mặt thực tế đã thu."),
                    main("Chủ xe kiểm tra thời gian nhận xe và tình trạng xe trước khi bàn giao."),
                    branch("Nếu danh tính, hợp đồng, thanh toán hoặc thời gian chưa hợp lệ, hệ thống từ chối bàn giao và giữ nguyên trạng thái booking."),
                    main("Nếu đủ điều kiện, chủ xe xác nhận giao xe và hệ thống ghi thời điểm bàn giao."),
                    main("Booking chuyển sang trạng thái đang thuê (IN_PROGRESS), xe được đánh dấu đang được thuê và các bên nhận thông báo."),
                ],
            ),
            proc(
                "Quy trình nhận xe trả và kiểm tra sau thuê",
                "Quy trình ghi nhận đầy đủ tình trạng xe tại thời điểm kết thúc chuyến thuê. Dữ liệu kiểm tra được so sánh với thời điểm bàn giao để xác định xe có thể hoàn tất ngay hay cần xử lý khoản phát sinh.",
                [
                    main("Người thuê trả xe tại địa điểm và thời gian được thống nhất."),
                    main("Chủ xe tiếp nhận xe của booking đang IN_PROGRESS và thuộc quyền quản lý."),
                    main("Hệ thống ghi nhận thời điểm trả xe thực tế."),
                    main("Chủ xe nhập số kilomet và mức nhiên liệu tại thời điểm nhận lại xe."),
                    main("Chủ xe chụp hoặc tải hình ảnh xe khi trả và mô tả tình trạng thực tế."),
                    main("Hệ thống kiểm tra dữ liệu kilomet, nhiên liệu, hình ảnh và nội dung tình trạng có đầy đủ, hợp lệ hay không."),
                    branch("Nếu dữ liệu chưa đầy đủ, hệ thống chưa cho kết thúc kiểm tra và yêu cầu bổ sung."),
                    main("Chủ xe so sánh tình trạng xe trả với tình trạng tại thời điểm bàn giao."),
                    main("Hệ thống lưu biên bản và chuyển booking sang giai đoạn kiểm tra sau thuê (RETURN_INSPECTION)."),
                    branch("Nếu không có phát sinh, biên bản có thể được xác nhận đã xử lý; nếu có phát sinh, booking chuyển sang bước xử lý phụ phí."),
                ],
            ),
            proc(
                "Quy trình xử lý phụ phí và hoàn tất chuyến thuê",
                "Quy trình gộp việc tạo, thanh toán, hủy phụ phí và kiểm tra điều kiện hoàn tất booking. Phụ phí phải dựa trên biên bản sau thuê và không được tạo ngoài giai đoạn kiểm tra phù hợp.",
                [
                    main("Chủ xe xác định khoản phát sinh từ biên bản kiểm tra sau thuê."),
                    main("Chủ xe chọn loại phụ phí, nhập số tiền, mô tả lý do và thêm minh chứng nếu có."),
                    main("Hệ thống kiểm tra booking đang ở giai đoạn kiểm tra, khoản tiền hợp lệ và chủ xe đúng quyền."),
                    branch("Nếu thiếu lý do, sai số tiền hoặc sai giai đoạn, hệ thống từ chối tạo phụ phí."),
                    main("Hệ thống tạo phụ phí gắn với booking và chuyển booking sang chờ xử lý phụ phí (AWAITING_EXTRA_CHARGE)."),
                    main("Người thuê xem khoản phí và thực hiện thanh toán bằng phương thức được hỗ trợ."),
                    main("Hệ thống tạo payment loại phụ phí và chỉ ghi nhận đã trả khi payment đạt PAID."),
                    main("Chủ xe có thể xác nhận tiền mặt thực tế đã thu hoặc hủy khoản phí không còn hợp lệ."),
                    main("Hệ thống kiểm tra tất cả phụ phí đã được thanh toán hoặc hủy, biên bản đã xử lý và không còn khoản thuê bắt buộc."),
                    branch("Nếu còn kiểm tra, remainingAmount hoặc phụ phí chờ, hệ thống chưa cho hoàn tất booking."),
                    main("Nếu mọi nghĩa vụ đã xử lý, chủ xe xác nhận hoàn tất chuyến thuê."),
                    main("Booking chuyển sang trạng thái hoàn tất (COMPLETED), xe trở lại trạng thái phù hợp và hợp đồng được đồng bộ."),
                    main("Hệ thống cập nhật lịch sử, Task Center và thông báo kết quả cho các bên."),
                ],
            ),
        ],
    },
    {
        "title": "Nghiệp vụ hợp đồng, lịch sử và đánh giá",
        "intro": [
            "Nhóm nghiệp vụ này lưu giữ căn cứ giao dịch và lịch sử sau khi booking được hình thành. Hợp đồng gắn với booking, lưu các bản chụp cần thiết và có thể được tạo sau khi chủ xe duyệt, trước hoặc trong quá trình người thuê thực hiện thanh toán.",
            "Lịch sử được lọc theo vai trò và quyền sở hữu. Sau khi booking COMPLETED, người thuê được đánh giá một lần; chủ xe có thể phản hồi, còn ADMIN xử lý nội dung bị báo cáo theo quyền kiểm duyệt.",
        ],
        "processes": [
            proc(
                "Quy trình tạo và xem hợp đồng",
                "Hợp đồng có thể được tạo ngay sau khi booking được chủ xe duyệt, không bắt buộc phải chờ thanh toán hoàn tất. Tuy nhiên, hợp đồng và nghĩa vụ thanh toán đều phải được hoàn thiện trước khi xe được bàn giao.",
                [
                    main("Booking chuyển sang OWNER_APPROVED sau khi chủ xe chấp nhận yêu cầu thuê."),
                    main("Người thuê yêu cầu tạo hoặc hoàn thiện hợp đồng trong giai đoạn trước hoặc trong quá trình thanh toán."),
                    main("Hệ thống kiểm tra booking thuộc người thuê, đã được duyệt và chưa ở trạng thái hủy, hoàn tất hoặc NO_SHOW."),
                    branch("Nếu booking không đúng quyền hoặc trạng thái, hệ thống từ chối tạo hợp đồng và không làm thay đổi booking."),
                    branch("Nếu hợp lệ, hệ thống tiếp tục chuẩn bị nội dung hợp đồng từ dữ liệu đã chốt của booking."),
                    main("Hệ thống chuẩn bị snapshot người thuê, chủ xe, xe, thời gian, địa điểm, giá, tiền cọc và số tiền còn lại."),
                    main("Hệ thống tạo mã hợp đồng và bảo đảm mỗi booking không phát sinh hợp đồng trùng."),
                    main("Người thuê và chủ xe xem hợp đồng theo quyền; người không liên quan không được truy cập."),
                    main("Hệ thống cập nhật trạng thái hợp đồng theo tiến trình nghiệp vụ."),
                    main("Trước khi bàn giao, hệ thống kiểm tra hợp đồng và nghĩa vụ thanh toán đã đáp ứng điều kiện nhận xe."),
                ],
            ),
            proc(
                "Quy trình xem lịch sử booking và thanh toán",
                "Lịch sử cung cấp dòng thời gian giao dịch theo đúng phạm vi của từng người sử dụng. Các thao tác lọc, tìm kiếm và mở chi tiết chỉ là bước giao diện trong quy trình này, không được tách thành nghiệp vụ riêng.",
                [
                    main("Người dùng mở lịch sử booking và thanh toán."),
                    main("Hệ thống xác định vai trò, danh tính và phạm vi sở hữu dữ liệu."),
                    branch("Người thuê chỉ xem booking và payment của mình."),
                    branch("Chủ xe chỉ xem booking liên quan đến xe mình; BUSINESS chỉ xem dữ liệu thuộc doanh nghiệp."),
                    main("Hệ thống tổng hợp trạng thái booking, payment, Refund, hợp đồng và phụ phí có liên quan."),
                    main("Người dùng lọc danh sách hoặc mở chi tiết một giao dịch trong phạm vi được phép."),
                    main("Hệ thống che hoặc không trả dữ liệu nhạy cảm của người khác."),
                    branch("Nếu yêu cầu vượt quyền, hệ thống từ chối truy cập và không tiết lộ dữ liệu."),
                ],
            ),
            proc(
                "Quy trình đánh giá, phản hồi và xử lý vi phạm",
                "Quy trình hợp nhất việc người thuê đánh giá, chủ xe phản hồi và ADMIN xử lý nội dung vi phạm. Đánh giá chỉ phát sinh sau chuyến thuê hoàn tất và không được tạo lặp cho cùng một booking.",
                [
                    main("Người thuê mở booking thuộc mình đã COMPLETED và chọn thực hiện đánh giá."),
                    main("Người thuê nhập điểm, nhận xét và hình ảnh nếu chức năng hình ảnh được sử dụng."),
                    main("Hệ thống kiểm tra booking đúng người thuê, đã hoàn tất và chưa có review của người thuê."),
                    branch("Nếu booking chưa đủ điều kiện hoặc đã có review, hệ thống từ chối tạo trùng."),
                    main("Hệ thống tạo review gắn với booking và xe, đồng thời chỉ công khai dữ liệu phù hợp."),
                    main("Chủ xe xem review thuộc xe mình và gửi một phản hồi theo quyền."),
                    main("Chủ xe hoặc người có quyền báo cáo nội dung có dấu hiệu vi phạm."),
                    main("ADMIN xem báo cáo, kiểm tra nội dung và quyết định giữ, ẩn hoặc hiển thị lại review."),
                    branch("Nếu báo cáo không có căn cứ, review giữ trạng thái hiện tại; nếu vi phạm, hệ thống áp dụng trạng thái kiểm duyệt phù hợp."),
                    main("Hệ thống không công khai dữ liệu nhạy cảm và thông báo kết quả xử lý cho các bên liên quan."),
                ],
            ),
        ],
    },
    {
        "title": "Nghiệp vụ thông báo và theo dõi hệ thống",
        "intro": [
            "Nhóm nghiệp vụ này hỗ trợ người sử dụng theo dõi sự kiện và công việc phát sinh từ các quy trình chính. Notification Center tập hợp thông báo theo người nhận; Task Center phản ánh các việc cần làm dựa trên trạng thái booking, payment, Refund, kiểm duyệt và các đối tượng liên quan.",
            "Dashboard tổng hợp dữ liệu theo vai trò mà không thay thế các nghiệp vụ nguồn. ADMIN xem phạm vi toàn hệ thống, BUSINESS và USER chủ xe ký gửi xem dữ liệu thuộc sở hữu, còn USER người thuê chỉ xem dữ liệu cá nhân khi màn hình tương ứng hỗ trợ.",
        ],
        "processes": [
            proc(
                "Quy trình xem thông báo và việc cần làm",
                "Thông báo không được tạo thành một nghiệp vụ độc lập cho từng sự kiện; nó là kết quả của quy trình đã phát sinh sự kiện. Quy trình này tập trung vào việc người nhận đọc, quản lý thông báo và điều hướng đến công việc đang chờ.",
                [
                    main("Hệ thống tạo thông báo khi quy trình nghiệp vụ phát sinh sự kiện cần gửi đến một vai trò hoặc người nhận cụ thể."),
                    main("Người dùng mở Notification Center để xem danh sách và trạng thái chưa đọc."),
                    main("Hệ thống chỉ trả thông báo thuộc đúng người nhận và phạm vi dữ liệu liên quan."),
                    main("Người dùng mở chi tiết; hệ thống đánh dấu đã đọc và điều hướng đến booking, payment, Refund hoặc đối tượng tương ứng."),
                    main("Người dùng có thể đánh dấu một thông báo hoặc tất cả thông báo là đã đọc."),
                    main("Khi người dùng xóa thông báo, hệ thống thực hiện xóa mềm trong phạm vi tài khoản."),
                    main("Người dùng mở Task Center để xem các việc cần làm được suy ra từ trạng thái nghiệp vụ."),
                    main("Khi trạng thái nguồn thay đổi, hệ thống cập nhật hoặc loại bỏ task tương ứng thay vì tạo một quy trình mới."),
                    branch("Nếu người dùng không có quyền với đối tượng đích, hệ thống không cho điều hướng hoặc xem dữ liệu chi tiết."),
                ],
            ),
            proc(
                "Quy trình xem dashboard theo vai trò",
                "Dashboard cung cấp số liệu tổng hợp nhằm hỗ trợ theo dõi vận hành. Mọi chỉ số phải được lọc theo quyền, trạng thái hợp lệ và phạm vi sở hữu trước khi hiển thị.",
                [
                    main("Người dùng mở dashboard và hệ thống xác định vai trò hiện tại."),
                    branch("ADMIN được tổng hợp xe, booking, payment, Refund, đánh giá và tài khoản trong phạm vi toàn hệ thống."),
                    branch("BUSINESS chỉ xem doanh nghiệp, xe, booking và dữ liệu tài chính thuộc doanh nghiệp."),
                    branch("USER chủ xe ký gửi chỉ xem xe và giao dịch thuộc sở hữu; USER người thuê xem dữ liệu cá nhân nếu dashboard hỗ trợ."),
                    main("Hệ thống đồng bộ các trạng thái hết hạn hoặc trạng thái xe cần cập nhật trước khi tổng hợp."),
                    main("Hệ thống chỉ tính payment PAID vào doanh thu và trừ các Refund SUCCEEDED để phản ánh doanh thu thuần theo logic dashboard hiện tại."),
                    main("Hệ thống tổng hợp số lượng xe, booking theo trạng thái, thanh toán, doanh thu, hoàn tiền và đánh giá trong phạm vi cho phép."),
                    branch("Nếu vai trò yêu cầu dữ liệu ngoài phạm vi, hệ thống loại dữ liệu đó khỏi kết quả."),
                    main("Dashboard hiển thị số liệu và liên kết đến danh sách chi tiết mà người dùng có quyền truy cập."),
                ],
            ),
        ],
    },
    {
        "title": "Quy trình thuê xe tổng thể",
        "intro": [
            "Quy trình thuê xe tổng thể liên kết bảy nhóm nghiệp vụ trước thành một vòng đời thống nhất từ tìm kiếm đến đánh giá. Người thuê có thể giữ xe trong giỏ hoặc tạo booking trực tiếp, sau đó chờ chủ xe quyết định và hoàn thiện hợp đồng cùng nghĩa vụ thanh toán.",
            "Khi đủ điều kiện, chủ xe bàn giao xe và booking bước vào giai đoạn đang thuê. Cuối chuyến, xe được tiếp nhận, kiểm tra và xử lý phụ phí nếu có; chỉ khi mọi nghĩa vụ đã hoàn tất thì booking mới chuyển COMPLETED.",
            "Các nhánh hủy, thanh toán thất bại, callback đến muộn, NO_SHOW và hoàn tiền chỉ được tóm tắt trong quy trình tổng thể. Nội dung chi tiết vẫn thuộc các quy trình chuyên trách để tránh lặp lại.",
        ],
        "processes": [
            proc(
                "Quy trình thuê xe tổng thể",
                "Quy trình thể hiện chuỗi nghiệp vụ chính và các điểm quyết định quan trọng của một chuyến thuê. Mỗi bước dẫn chiếu đến nhóm chuyên trách thay vì chép lại toàn bộ quy tắc chi tiết.",
                [
                    main("Người thuê tìm kiếm xe theo địa điểm, thời gian và tiêu chí; hệ thống chỉ trả xe hợp lệ và khả dụng."),
                    branch("Nếu xe hoặc khoảng thời gian không hợp lệ, người thuê phải điều chỉnh điều kiện tìm kiếm."),
                    main("Người thuê xem chi tiết, giá dự kiến, đánh giá và vị trí nhận xe trên bản đồ."),
                    main("Người thuê thêm xe vào giỏ để giữ tạm thời hoặc chọn tạo booking trực tiếp."),
                    branch("Nếu giỏ hết hạn, hệ thống giải phóng lượt giữ và yêu cầu kiểm tra khả dụng lại."),
                    main("Hệ thống kiểm tra quyền đặt, lịch xe, thông tin người thuê và tạo booking gắn đúng chủ xe."),
                    main("Chủ xe duyệt hoặc từ chối booking; booking bị từ chối không tiếp tục và vấn đề tài chính được chuyển sang quy trình chuyên trách."),
                    main("Khi được duyệt, người thuê tạo hoặc hoàn thiện hợp đồng và thực hiện thanh toán theo lựa chọn."),
                    branch("Nếu payment thất bại, nghĩa vụ chưa được ghi nhận; nếu callback đến sau khi booking đã hủy, booking không được kích hoạt lại."),
                    branch("Nếu người có quyền hủy booking, hệ thống xử lý CANCELLED và hoàn tiền thủ công theo quy trình 2.4.1.4.3."),
                    main("Chủ xe kiểm tra người thuê, hợp đồng, thanh toán và thời gian trước khi bàn giao xe."),
                    branch("Nếu người thuê không đến sau giờ nhận xe và đủ khoảng chờ 30 phút, chủ xe có thể ghi nhận NO_SHOW."),
                    main("Khi bàn giao thành công, booking chuyển IN_PROGRESS và chuyến thuê diễn ra."),
                    main("Người thuê trả xe; chủ xe ghi thời điểm, kilomet, nhiên liệu, hình ảnh và tình trạng xe."),
                    branch("Nếu xe có hư hỏng hoặc khoản phát sinh, chủ xe tạo phụ phí và người thuê xử lý payment liên quan."),
                    branch("Nếu còn remainingAmount, phụ phí hoặc kiểm tra chưa hoàn tất, hệ thống chưa cho kết thúc booking."),
                    main("Khi mọi nghĩa vụ đã xử lý, chủ xe hoàn tất booking; booking chuyển COMPLETED và xe trở lại trạng thái phù hợp."),
                    main("Người thuê thực hiện đánh giá; lịch sử, dashboard, Task Center và Notification Center được cập nhật theo quyền."),
                ],
            ),
        ],
    },
]


def set_font(run, size=13, bold=None, italic=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    rpr = run._element.get_or_add_rPr()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rpr.rFonts.set(qn(key), FONT)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clear_body(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype in {RT.IMAGE, RT.HYPERLINK}:
            doc.part.drop_rel(rid)


def clear_part(part):
    for paragraph in list(part.paragraphs):
        for child in list(paragraph._element):
            paragraph._element.remove(child)
    for rid, rel in list(part.part.rels.items()):
        if rel.reltype == RT.IMAGE:
            part.part.drop_rel(rid)


def configure_style(style, size, bold=False, before=0, after=3, keep=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = BLACK
    rpr = style._element.get_or_add_rPr()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rpr.rFonts.set(qn(key), FONT)
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = 1.5
    fmt.keep_with_next = keep
    fmt.widow_control = True


def add_field(paragraph, instruction, cached=""):
    run = paragraph.add_run()
    set_font(run)
    for kind, value in (("begin", None), (None, instruction), ("separate", None), (None, cached), ("end", None)):
        if kind:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        else:
            node = OxmlElement("w:instrText" if value == instruction else "w:t")
            node.set(qn("xml:space"), "preserve")
            node.text = value
        run._r.append(node)


def add_update_fields(doc):
    settings = doc.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


def add_page_number(section):
    clear_part(section.footer)
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    add_field(paragraph, " PAGE ", "1")


def prune_unused_headers_footers(doc):
    used = {
        node.get(qn("r:id"))
        for node in doc.sections[0]._sectPr
        if node.tag in {qn("w:headerReference"), qn("w:footerReference")}
    }
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype in {RT.HEADER, RT.FOOTER} and rid not in used:
            doc.part.drop_rel(rid)


def add_bullet_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(e.get(qn("w:abstractNumId"))) for e in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(e.get(qn("w:numId"))) for e in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.append(multi)
    for level, marker, left in ((0, "●", 720), (1, "○", 1440)):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        for tag, value in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", marker), ("w:suff", "space")):
            node = OxmlElement(tag)
            node.set(qn("w:val"), value)
            lvl.append(node)
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(left))
        ind.set(qn("w:hanging"), "360")
        ppr.append(ind)
        lvl.append(ppr)
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), FONT)
        fonts.set(qn("w:hAnsi"), FONT)
        rpr.append(fonts)
        lvl.append(rpr)
        abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def add_step(doc, text, level, bullet_num, step_number=None):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.keep_together = True
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(bullet_num))
    numpr.extend([ilvl, numid])
    ppr.append(numpr)
    prefix = f"Bước {step_number}: " if level == 0 else ""
    run = paragraph.add_run(prefix + text)
    set_font(run)
    return paragraph


def build():
    assert len(GROUPS) == 8
    assert [len(group["processes"]) for group in GROUPS] == [3, 3, 5, 3, 3, 3, 2, 1]
    assert sum(len(group["processes"]) for group in GROUPS) == 23

    doc = Document(REFERENCE)
    clear_body(doc)
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    section.different_first_page_header_footer = False
    clear_part(section.header)
    add_page_number(section)
    prune_unused_headers_footers(doc)

    configure_style(doc.styles["Normal"], 13, after=3)
    doc.styles["Normal"].paragraph_format.first_line_indent = Cm(1)
    doc.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    configure_style(doc.styles["Heading 1"], 14, bold=True, before=10, after=6, keep=True)
    configure_style(doc.styles["Heading 2"], 13, bold=True, before=8, after=3, keep=True)
    configure_style(doc.styles["Heading 3"], 13, bold=True, before=6, after=2, keep=True)
    if "Caption" not in [style.name for style in doc.styles]:
        doc.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    add_update_fields(doc)
    bullet_num = add_bullet_numbering(doc)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(115)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    run = title.add_run("CÁC NGHIỆP VỤ RÚT GỌN\nCỦA HỆ THỐNG BQDRIVE")
    set_font(run, size=16, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.first_line_indent = Cm(0)
    set_font(subtitle.add_run("Tổ chức theo 8 nhóm và 23 quy trình chính"), italic=True)
    doc.add_page_break()

    toc_title = doc.add_paragraph()
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_title.paragraph_format.first_line_indent = Cm(0)
    set_font(toc_title.add_run("MỤC LỤC"), size=14, bold=True)
    toc = doc.add_paragraph()
    toc.paragraph_format.first_line_indent = Cm(0)
    add_field(toc, ' TOC \\o "1-2" \\h \\z \\u ', "Mục lục được cập nhật khi mở tài liệu.")
    doc.add_page_break()

    opening = [
        "Các nghiệp vụ của BQDrive được tổ chức lại thành tám nhóm lớn và hai mươi ba quy trình chính nhằm giảm thao tác lặp, nhưng vẫn giữ nguyên vai trò, trạng thái, quyền hạn, điều kiện tài chính và các quy tắc đã đối chiếu từ hệ thống.",
        "Mỗi nhóm được giới thiệu bằng văn xuôi, sau đó trình bày các quy trình bằng các bước tuần tự và nhánh xử lý. Các thao tác tra cứu, CRUD, loại payment và hành động thông báo được đặt trong quy trình chuyên trách thay vì tách thành các quy trình nhỏ độc lập.",
    ]
    for text in opening:
        p = doc.add_paragraph()
        set_font(p.add_run(text))

    process_count = 0
    for group_index, group in enumerate(GROUPS, 1):
        group_heading = doc.add_paragraph(
            f"2.4.1.{group_index} {group['title']}",
            style="Heading 1",
        )
        for run in group_heading.runs:
            set_font(run, size=14, bold=True)
        for text in group["intro"]:
            paragraph = doc.add_paragraph()
            set_font(paragraph.add_run(text))
        for process_index, process in enumerate(group["processes"], 1):
            process_count += 1
            heading = doc.add_paragraph(
                f"2.4.1.{group_index}.{process_index} {process['title']}",
                style="Heading 2",
            )
            for run in heading.runs:
                set_font(run, bold=True)
            lead = doc.add_paragraph()
            set_font(lead.add_run(process["lead"]))
            label = doc.add_paragraph()
            label.paragraph_format.first_line_indent = Cm(0)
            label.paragraph_format.keep_with_next = True
            set_font(label.add_run("Các bước thực hiện:"), bold=True)
            main_step_number = 0
            for level, text in process["steps"]:
                if level == 0:
                    main_step_number += 1
                    add_step(doc, text, level, bullet_num, main_step_number)
                else:
                    add_step(doc, text, level, bullet_num)

    assert process_count == 23
    doc.core_properties.title = "Các nghiệp vụ rút gọn của hệ thống BQDrive"
    doc.core_properties.subject = "8 nhóm nghiệp vụ và 23 quy trình chính"
    doc.core_properties.author = "BQDrive"
    doc.save(OUTPUT)

    summary = {
        "output": str(OUTPUT),
        "groups": [group["title"] for group in GROUPS],
        "group_process_counts": [len(group["processes"]) for group in GROUPS],
        "process_count": process_count,
        "diagram_count": 0,
        "processes": [
            f"2.4.1.{group_index}.{process_index} {process['title']}"
            for group_index, group in enumerate(GROUPS, 1)
            for process_index, process in enumerate(group["processes"], 1)
        ],
    }
    (WORK / "reduced_build_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build()
