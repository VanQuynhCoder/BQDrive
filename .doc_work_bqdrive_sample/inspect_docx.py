from pathlib import Path
import hashlib
import json
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn


def run_font_summary(paragraph):
    result = []
    for run in paragraph.runs:
        if not run.text.strip():
            continue
        result.append(
            {
                "text": run.text[:80],
                "font": run.font.name,
                "size": run.font.size.pt if run.font.size else None,
                "bold": run.bold,
                "italic": run.italic,
            }
        )
    return result


def main():
    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    document = Document(source)
    records = []
    page = 1
    for index, paragraph in enumerate(document.paragraphs):
        text = " ".join(paragraph.text.split())
        has_drawing = bool(paragraph._p.xpath(".//w:drawing"))
        record = {
            "index": index,
            "page": page,
            "style": paragraph.style.name if paragraph.style else "",
            "text": text,
            "alignment": int(paragraph.alignment) if paragraph.alignment is not None else None,
            "left_indent_cm": (
                paragraph.paragraph_format.left_indent.cm
                if paragraph.paragraph_format.left_indent is not None
                else None
            ),
            "first_line_cm": (
                paragraph.paragraph_format.first_line_indent.cm
                if paragraph.paragraph_format.first_line_indent is not None
                else None
            ),
            "space_before_pt": (
                paragraph.paragraph_format.space_before.pt
                if paragraph.paragraph_format.space_before is not None
                else None
            ),
            "space_after_pt": (
                paragraph.paragraph_format.space_after.pt
                if paragraph.paragraph_format.space_after is not None
                else None
            ),
            "line_spacing": paragraph.paragraph_format.line_spacing,
            "drawing": has_drawing,
            "runs": run_font_summary(paragraph),
        }
        if text or has_drawing:
            records.append(record)
        page_breaks = len(paragraph._p.xpath(".//w:lastRenderedPageBreak"))
        explicit_breaks = len(paragraph._p.xpath('.//w:br[@w:type="page"]'))
        page += page_breaks + explicit_breaks

    with zipfile.ZipFile(source) as archive:
        package = [
            {
                "path": info.filename,
                "size": info.file_size,
                "sha256": hashlib.sha256(archive.read(info.filename)).hexdigest(),
            }
            for info in archive.infolist()
        ]
    payload = {
        "source": str(source),
        "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "section_count": len(document.sections),
        "estimated_saved_pages": page,
        "records": records,
        "package": package,
    }
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {key: payload[key] for key in payload if key not in {"records", "package"}},
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
