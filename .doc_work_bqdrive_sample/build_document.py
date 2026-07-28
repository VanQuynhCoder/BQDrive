from pathlib import Path
import json
import math
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"D:\DoAn_LuanVanTotNghiep")
WORK = ROOT / ".doc_work_bqdrive_sample"
REFERENCE = Path(r"C:\Users\bom13\Downloads\VuHoangUng-1.docx")
OUTPUT = ROOT / "BQDrive_Cac_nghiep_vu_trinh_bay_theo_mau.docx"
DIAGRAM_DIR = WORK / "diagrams"
DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)


def P(title, actor, action, check, success, state="", invalid="Dữ liệu hoặc điều kiện không hợp lệ; hệ thống từ chối xử lý và nêu lý do.", diagram=None, extra=None):
    return {
        "title": title, "actor": actor, "action": action, "check": check,
        "success": success, "state": state, "invalid": invalid,
        "diagram": diagram, "extra": extra or [],
    }


GROUPS = [
    {
        "title": "Nghiệp vụ tài khoản và phân quyền",
        "intro": [
            "Nghiệp vụ tài khoản và phân quyền bảo đảm mỗi người tham gia BQDrive có danh tính hợp lệ và chỉ sử dụng các chức năng phù hợp. USER có thể tự đăng ký, xác thực email và sử dụng hệ thống với vai trò người thuê xe hoặc chủ xe ký gửi; BUSINESS được ADMIN tạo hoặc cấp tài khoản sau khi thông tin được kiểm tra.",
            "Quy trình bao gồm đăng ký, xác thực OTP, đăng nhập bằng email hoặc Google OAuth, khôi phục mật khẩu, cập nhật hồ sơ và quản trị trạng thái tài khoản. Google OAuth có thể tạo một tài khoản USER mới từ email Google đã xác thực khi hệ thống chưa có tài khoản tương ứng.",
            "Email phải đúng định dạng và không trùng. OTP phải chính xác, còn hiệu lực và chưa vượt quá số lần thử. Tài khoản bị khóa hoặc xóa mềm không được đăng nhập; cập nhật hồ sơ không làm thay đổi vai trò và ADMIN không được cấp quyền trái chính sách.",
        ],
        "processes": [
            P("Quy trình đăng ký tài khoản USER", "Người đăng ký", "cung cấp họ tên, email, mật khẩu và thông tin bắt buộc", "email đúng định dạng, chưa tồn tại và dữ liệu bắt buộc hợp lệ", "tạo tài khoản USER và gửi OTP đến email", "Tài khoản ở trạng thái chờ xác thực email.", "Email trùng hoặc dữ liệu sai; người đăng ký phải điều chỉnh.", "register"),
            P("Quy trình xác thực email bằng OTP", "USER", "nhập OTP đã nhận qua email", "OTP thuộc đúng tài khoản, còn hiệu lực và khớp mã đã phát hành", "xác nhận email của tài khoản", "Tài khoản chuyển sang trạng thái đã xác thực.", "OTP sai, hết hạn hoặc vượt số lần thử; hệ thống không xác thực.", "verify_email"),
            P("Quy trình đăng nhập bằng email và mật khẩu", "Người dùng", "cung cấp email và mật khẩu", "thông tin xác thực đúng, tài khoản đã xác thực, không bị khóa và không bị xóa mềm", "tạo phiên đăng nhập theo vai trò ADMIN, BUSINESS hoặc USER", "", "Thông tin sai hoặc tài khoản không đủ điều kiện; hệ thống từ chối đăng nhập.", "login"),
            P("Quy trình đăng nhập bằng Google OAuth", "Người dùng", "chọn tài khoản Google và đồng ý xác thực", "thông tin Google hợp lệ, email có thể đối chiếu và tài khoản không bị vô hiệu hóa", "đăng nhập tài khoản hiện có hoặc tạo tài khoản USER từ email Google đã xác thực", "", "Google không xác thực được hoặc tài khoản đã bị khóa/xóa mềm; hệ thống từ chối đăng nhập."),
            P("Quy trình quên mật khẩu và đặt lại mật khẩu", "Người dùng", "yêu cầu khôi phục bằng email và nhập OTP", "email thuộc tài khoản hoạt động, OTP còn hạn và mã đặt lại hợp lệ", "cho phép thiết lập mật khẩu mới và vô hiệu hóa thông tin khôi phục cũ", "", "Email không tồn tại, OTP sai/hết hạn hoặc mã đặt lại không hợp lệ; mật khẩu không đổi."),
            P("Quy trình cập nhật hồ sơ cá nhân", "Người dùng", "chỉnh sửa các trường hồ sơ được phép", "dữ liệu mới hợp lệ và người dùng đang thao tác trên tài khoản của mình", "cập nhật hồ sơ nhưng giữ nguyên vai trò và quyền", "", "Dữ liệu sai hoặc yêu cầu thay đổi vai trò; hệ thống từ chối."),
            P("Quy trình ADMIN tạo hoặc cấp tài khoản BUSINESS", "ADMIN", "nhập thông tin doanh nghiệp và người đại diện cần cấp tài khoản", "email chưa trùng, thông tin doanh nghiệp hợp lệ và quyền cấp phù hợp chính sách", "tạo hồ sơ BUSINESS và tài khoản sử dụng tương ứng", "Tài khoản được gán vai trò BUSINESS.", "Thông tin thiếu, email trùng hoặc yêu cầu quyền không phù hợp; hệ thống không tạo.", "business_account"),
            P("Quy trình ADMIN khóa, mở khóa hoặc xóa mềm tài khoản", "ADMIN", "chọn tài khoản USER hoặc BUSINESS và trạng thái cần áp dụng", "tài khoản tồn tại và thao tác không vi phạm chính sách quản trị", "ghi nhận trạng thái khóa, mở khóa hoặc xóa mềm mà vẫn giữ lịch sử", "", "Tài khoản không tồn tại hoặc thao tác không được phép; trạng thái được giữ nguyên."),
        ],
    },
    {
        "title": "Nghiệp vụ quản lý và kiểm duyệt xe",
        "intro": [
            "Nghiệp vụ này cho phép chủ xe, gồm BUSINESS và USER ký gửi, quản lý xe thuộc quyền sở hữu của mình. Hồ sơ xe bao gồm biển số, thương hiệu, hình ảnh, thông số kỹ thuật, giá, hình thức thuê và vị trí nhận xe.",
            "Xe mới phải chờ ADMIN kiểm duyệt trước khi có thể được cung cấp cho người thuê. Thay đổi quan trọng trên xe đã duyệt có thể đưa xe trở lại trạng thái chờ duyệt; cập nhật riêng vị trí nhận xe được xử lý độc lập và có thể giữ trạng thái đã duyệt.",
            "Địa chỉ nhận xe, vĩ độ và kinh độ được lưu để phục vụ bản đồ và tìm kiếm theo địa điểm. Dữ liệu công khai ưu tiên xe đã duyệt; code hiện cũng cho phép một số truy vấn công khai giữ xe đang được thuê để thể hiện tình trạng, nhưng xe đó không được coi là khả dụng cho booking mới.",
        ],
        "processes": [
            P("Quy trình xem danh sách và chi tiết xe của chủ xe", "Chủ xe", "yêu cầu xem các xe thuộc phạm vi quản lý", "vai trò và quyền sở hữu của từng xe", "hiển thị hồ sơ, trạng thái kiểm duyệt, giá và vị trí nhận xe", "", "Xe không thuộc chủ xe; hệ thống không cung cấp chi tiết."),
            P("Quy trình thêm xe mới", "Chủ xe", "nhập biển số, thương hiệu, hình ảnh, thông số, giá và vị trí nhận xe", "quyền chủ xe, biển số không trùng, giá/hình thức thuê và địa chỉ hợp lệ", "tạo hồ sơ xe để ADMIN kiểm duyệt", "Xe chuyển sang trạng thái chờ duyệt (PENDING).", "Thiếu dữ liệu, biển số trùng hoặc vị trí không hợp lệ; xe không được tạo.", "add_car"),
            P("Quy trình cập nhật thông tin xe", "Chủ xe", "chỉnh sửa xe thuộc quyền quản lý", "dữ liệu mới hợp lệ và xác định thay đổi có thuộc nhóm quan trọng", "lưu thông tin mới", "Nếu thay đổi quan trọng trên xe đã duyệt, xe chuyển về chờ duyệt (PENDING).", "Xe không thuộc quyền hoặc dữ liệu không hợp lệ; hệ thống không cập nhật."),
            P("Quy trình ẩn, hiển thị lại và xóa mềm xe", "Chủ xe", "chọn trạng thái hiển thị hoặc xóa mềm cho xe", "quyền sở hữu và tình trạng hiện tại của xe", "ẩn xe, hiển thị lại xe đủ điều kiện hoặc đánh dấu xóa mềm", "", "Xe không thuộc quyền hoặc không đủ điều kiện hiển thị lại; thao tác bị từ chối."),
            P("Quy trình quản lý thương hiệu xe", "ADMIN", "thêm, sửa hoặc xóa mềm thương hiệu dùng chung", "tên thương hiệu hợp lệ, không trùng và thao tác không phá vỡ dữ liệu đang dùng", "cập nhật danh mục thương hiệu phục vụ hồ sơ xe", "", "Thương hiệu trùng hoặc dữ liệu không hợp lệ; danh mục giữ nguyên."),
            P("Quy trình quản lý vị trí nhận xe trên bản đồ", "Chủ xe", "chọn điểm nhận xe và cung cấp địa chỉ hoặc ghi chú", "vĩ độ từ -90 đến 90, kinh độ từ -180 đến 180 và xe thuộc quyền quản lý", "lưu địa chỉ, tọa độ, lịch sử cập nhật vị trí và người cập nhật", "Nếu chỉ thay đổi vị trí, xe đã duyệt có thể giữ trạng thái đã duyệt (APPROVED).", "Tọa độ sai hoặc xe không thuộc quyền; hệ thống không cập nhật vị trí.", "location"),
            P("Quy trình ADMIN duyệt xe", "ADMIN", "kiểm tra xe đang chờ duyệt", "hình ảnh, biển số, kỹ thuật, giá, thương hiệu và vị trí đáp ứng yêu cầu", "cho phép xe tham gia danh sách xe hợp lệ", "Xe chuyển sang trạng thái đã duyệt (APPROVED).", "Xe chưa đạt yêu cầu; ADMIN không duyệt và phải xử lý theo quy trình từ chối.", "car_approval"),
            P("Quy trình ADMIN từ chối xe", "ADMIN", "chọn từ chối xe sau khi kiểm tra", "xe đang chờ duyệt và lý do từ chối đã được cung cấp", "lưu lý do để chủ xe chỉnh sửa", "Xe chuyển sang trạng thái bị từ chối (REJECTED).", "Thiếu lý do hoặc xe không còn chờ duyệt; hệ thống từ chối thao tác."),
            P("Quy trình cập nhật xe sau khi đã được duyệt", "Chủ xe", "cập nhật hồ sơ xe đang ở trạng thái đã duyệt", "so sánh dữ liệu cũ và mới để xác định thay đổi quan trọng", "lưu thay đổi và thông báo kết quả kiểm duyệt lại", "Thay đổi quan trọng đưa xe về PENDING; thay đổi không quan trọng giữ APPROVED.", "Dữ liệu sai hoặc xe không thuộc quyền; hệ thống giữ nguyên hồ sơ."),
        ],
    },
    {
        "title": "Nghiệp vụ tìm kiếm, giỏ xe và vòng đời booking",
        "intro": [
            "Người thuê xe bắt đầu bằng việc tìm kiếm theo địa điểm, thời gian và các tiêu chí như thương hiệu, giá, số ghế, nhiên liệu hoặc hộp số. Sau khi xem chi tiết và vị trí nhận xe, người thuê có thể giữ xe tạm thời trong giỏ hoặc tạo booking trực tiếp.",
            "Hệ thống kiểm tra xe, quyền đặt và lịch thuê trước khi giữ hoặc tạo booking. Người thuê không được đặt xe của mình; thời gian bắt đầu phải trước thời gian kết thúc và không được trùng với booking đang giữ lịch.",
            "Booking mới chờ chủ xe duyệt (REQUESTED). Chủ xe chỉ xử lý booking của xe thuộc quyền quản lý và có thể duyệt, từ chối kèm lý do hoặc hủy theo trạng thái cho phép. Nếu người thuê không đến sau giờ nhận xe và khoảng chờ 30 phút, chủ xe có thể xử lý khách không nhận xe khi xe chưa bàn giao.",
        ],
        "processes": [
            P("Quy trình tìm kiếm và lọc xe", "Người thuê xe", "nhập địa điểm, thời gian và tiêu chí xe", "khoảng thời gian hợp lệ và tiêu chí có thể áp dụng", "trả về xe phù hợp, kèm trạng thái khả dụng", "", "Không có xe phù hợp; hệ thống yêu cầu điều chỉnh tiêu chí.", "search"),
            P("Quy trình xem chi tiết xe", "Người thuê xe", "chọn một xe từ kết quả tìm kiếm", "xe tồn tại, không bị xóa mềm và được phép hiển thị", "hiển thị hình ảnh, kỹ thuật, giá, điều kiện thuê và vị trí bản đồ", "", "Xe không còn hợp lệ; hệ thống không cho tiếp tục đặt."),
            P("Quy trình thêm xe vào giỏ và giữ xe tạm thời", "Người thuê xe", "chọn xe và khoảng thời gian muốn giữ", "xe đã duyệt, không bị ẩn/xóa, không thuộc người thuê và không trùng lịch", "tạo lượt giữ có thời hạn và lưu bản chụp giá", "Lượt giữ ở trạng thái đang hoạt động (ACTIVE).", "Xe hoặc thời gian không hợp lệ; hệ thống không giữ xe.", "cart_hold"),
            P("Quy trình người thuê hủy lượt giữ xe", "Người thuê xe", "chọn hủy lượt giữ của mình", "lượt giữ thuộc đúng người thuê, chưa tạo booking và còn có thể hủy", "giải phóng khoảng thời gian giữ xe", "Lượt giữ chuyển sang đã hủy (CANCELLED).", "Lượt giữ không thuộc người thuê hoặc đã được sử dụng; hệ thống từ chối."),
            P("Quy trình hệ thống xử lý lượt giữ hết hạn", "Hệ thống", "đối chiếu thời điểm hết hạn của các lượt giữ", "lượt giữ vẫn ACTIVE và đã quá thời hạn", "giải phóng xe để người khác có thể chọn", "Lượt giữ chuyển sang hết hạn (EXPIRED).", "Lượt giữ đã tạo booking hoặc chưa hết hạn; hệ thống giữ nguyên."),
            P("Quy trình tạo booking trực tiếp", "Người thuê xe", "gửi yêu cầu thuê từ trang chi tiết xe", "xe, quyền đặt, thời gian, lịch thuê, thông tin người thuê và bản chụp giá hợp lệ", "tạo booking gắn đúng người thuê, xe và chủ xe", "Booking chuyển sang chờ chủ xe duyệt (REQUESTED).", "Một điều kiện không đạt; hệ thống không tạo booking.", "create_booking"),
            P("Quy trình tạo booking từ giỏ", "Người thuê xe", "chuyển lượt giữ trong giỏ thành booking", "lượt giữ thuộc người thuê, còn hiệu lực và xe vẫn khả dụng", "tạo booking từ dữ liệu và bản chụp giá của lượt giữ", "Booking chuyển sang chờ chủ xe duyệt (REQUESTED); lượt giữ chuyển sang đã đặt (BOOKED).", "Lượt giữ hết hạn hoặc không thuộc người thuê; hệ thống từ chối."),
            P("Quy trình chủ xe duyệt booking", "Chủ xe", "xem và chấp nhận yêu cầu thuê liên quan đến xe của mình", "booking đang chờ chủ xe duyệt (REQUESTED), xe thuộc quyền quản lý và lịch vẫn hợp lệ", "xác nhận booking để người thuê tạo hợp đồng và thanh toán", "Booking chuyển sang chủ xe đã duyệt (OWNER_APPROVED).", "Booking không thuộc quyền hoặc không còn hợp lệ; hệ thống không duyệt.", "approve_booking"),
            P("Quy trình chủ xe từ chối booking", "Chủ xe", "từ chối yêu cầu thuê và nhập lý do", "booking thuộc xe của chủ xe và đang chờ duyệt", "lưu lý do từ chối và giải phóng lịch", "Booking chuyển sang bị từ chối (REJECTED).", "Thiếu lý do hoặc booking không đúng trạng thái; hệ thống từ chối thao tác."),
            P("Quy trình hủy booking", "Người thuê xe hoặc Chủ xe", "xem trước chính sách, cung cấp lý do và xác nhận hủy", "quyền hủy và trạng thái thuộc chờ chủ xe duyệt (REQUESTED), chủ xe đã duyệt (OWNER_APPROVED), chờ thanh toán (PAYMENT_PENDING), đã thanh toán (PAID) hoặc trạng thái cũ tương ứng", "hủy booking, giải phóng lịch và tính khoản hoàn nếu có", "Booking chuyển sang đã hủy (CANCELLED).", "Booking đã ở trạng thái đang thuê (IN_PROGRESS), kiểm tra sau thuê (RETURN_INSPECTION), chờ xử lý phụ phí (AWAITING_EXTRA_CHARGE), hoàn tất (COMPLETED), bị từ chối (REJECTED) hoặc khách không nhận xe (NO_SHOW); không hủy theo luồng thường.", "cancel_refund"),
            P("Quy trình xử lý khách không đến nhận xe", "Chủ xe", "xác nhận người thuê không đến nhận xe và cung cấp lý do", "đã quá giờ nhận xe cùng khoảng chờ 30 phút, xe chưa bàn giao và booking thuộc trạng thái cho phép", "ghi nhận no-show và giải phóng lịch xe", "Booking chuyển sang khách không nhận xe (NO_SHOW).", "Chưa hết thời gian chờ, xe đã bàn giao hoặc trạng thái không hợp lệ; hệ thống từ chối.", "no_show"),
        ],
    },
    {
        "title": "Nghiệp vụ tính giá, thanh toán và hoàn tiền",
        "intro": [
            "Giá thuê được xác định theo hình thức thuê theo giờ hoặc theo ngày. Với thuê theo ngày, hệ thống phân loại từng ngày theo thứ tự ưu tiên ngày lễ, cuối tuần rồi ngày thường; ADMIN duy trì lịch ngày lễ làm căn cứ tính giá.",
            "Bản chụp giá được lưu khi giữ xe hoặc tạo booking để totalPrice không thay đổi khi chủ xe cập nhật giá sau đó. Sau khi chủ xe duyệt booking, người thuê có thể thanh toán cọc, toàn bộ hoặc phần còn lại bằng tiền mặt, MoMo hoặc VNPay.",
            "Chỉ payment đã thanh toán thành công (PAID) được cộng. Thanh toán cọc không phải thanh toán đủ; số tiền còn lại phải bằng 0 mới được xem là hoàn tất nghĩa vụ. Khi booking hủy hợp lệ, hệ thống tính chính sách và tạo Refund nếu số tiền hoàn lớn hơn 0; code hiện theo dõi việc chủ xe hoàn thủ công và người thuê xác nhận.",
        ],
        "processes": [
            P("Quy trình ADMIN quản lý ngày lễ", "ADMIN", "xem, thêm, cập nhật hoặc tắt ngày lễ", "thời gian áp dụng hợp lệ và không tạo dữ liệu trùng không cần thiết", "cập nhật lịch ngày lễ dùng cho tính giá", "", "Dữ liệu ngày lễ sai hoặc trùng; hệ thống không áp dụng."),
            P("Quy trình tính giá thuê theo ngày", "Hệ thống", "phân loại từng ngày trong khoảng thuê", "ngày thuê thuộc ngày lễ, cuối tuần hoặc ngày thường theo đúng thứ tự ưu tiên", "cộng giá từng ngày thành totalPrice", "", "Thiếu mức giá cần thiết hoặc khoảng thời gian sai; hệ thống không xác nhận giá."),
            P("Quy trình tính giá thuê theo giờ", "Hệ thống", "tính số giờ trong khoảng thuê", "xe hỗ trợ thuê theo giờ, thời lượng và giá theo giờ hợp lệ", "nhân số giờ với đơn giá và tạo tổng giá", "", "Xe không hỗ trợ hoặc thời lượng không hợp lệ; hệ thống từ chối."),
            P("Quy trình lưu bản chụp giá tại thời điểm đặt", "Hệ thống", "ghi nhận mức giá khi giữ xe hoặc tạo booking", "kết quả tính giá và thông tin xe hợp lệ", "lưu giá ngày thường, cuối tuần, ngày lễ hoặc giá giờ cùng totalPrice", "", "Không xác định được giá; hệ thống không cho hoàn tất việc giữ hoặc đặt."),
            P("Quy trình lựa chọn loại và phương thức thanh toán", "Người thuê xe", "chọn DEPOSIT, FULL hoặc REMAINING và CASH, MoMo hoặc VNPay", "booking đã được chủ xe duyệt và lựa chọn phù hợp số tiền còn lại", "tạo payment theo loại và phương thức đã chọn", "Payment ban đầu chờ xử lý (PENDING), trừ trường hợp tiền mặt được xác nhận ngay.", "Loại, phương thức hoặc số tiền không hợp lệ; payment không được tạo.", "payment"),
            P("Quy trình thanh toán cọc", "Người thuê xe", "thực hiện thanh toán khoản cọc", "booking cho phép thanh toán và chưa có khoản cọc PAID bị ghi nhận trùng", "ghi nhận payment cọc khi giao dịch thành công", "Booking có thể chuyển sang đã thanh toán (PAID) nhưng vẫn còn số tiền phải trả.", "Thanh toán thất bại hoặc trùng giao dịch; khoản tiền không được cộng."),
            P("Quy trình thanh toán toàn bộ", "Người thuê xe", "thanh toán toàn bộ totalPrice", "chưa có khoản đã trả gây vượt tổng tiền và giao dịch hợp lệ", "ghi nhận payment FULL thành công", "Số tiền còn lại bằng 0 và booking được ghi nhận thanh toán đủ.", "Giao dịch thất bại, sai số tiền hoặc đã xử lý; hệ thống không cộng."),
            P("Quy trình thanh toán phần còn lại", "Người thuê xe", "thanh toán số tiền còn lại sau khoản cọc", "remainingAmount lớn hơn 0 và booking thuộc trạng thái cho phép", "ghi nhận payment REMAINING", "Khi remainingAmount bằng 0, booking được xem là thanh toán đủ.", "Không còn số tiền phải trả hoặc giao dịch sai; hệ thống từ chối."),
            P("Quy trình xử lý kết quả thanh toán MoMo hoặc VNPay", "Hệ thống", "tiếp nhận kết quả giao dịch từ cổng thanh toán", "mã giao dịch, payment, booking, số tiền và dữ liệu xác thực khớp, chưa xử lý trước đó", "cập nhật payment đúng một lần", "Payment thành công chuyển sang đã thanh toán (PAID); payment không thành công chuyển sang thất bại (FAILED).", "Kết quả không hợp lệ hoặc trùng; hệ thống không ghi nhận lại."),
            P("Quy trình chủ xe xác nhận thanh toán tiền mặt", "Chủ xe", "xác nhận khoản tiền mặt thực tế đã thu", "booking thuộc xe của chủ xe, số tiền đúng và payment đang chờ", "ghi nhận payment tiền mặt đã thu", "Payment chuyển sang đã thanh toán (PAID).", "Chưa thu tiền, sai số tiền hoặc không đúng quyền; hệ thống không xác nhận."),
            P("Quy trình xác định chính sách và số tiền hoàn", "Hệ thống", "tổng hợp payment PAID của booking bị hủy", "booking hủy hợp lệ, loại payment thuộc DEPOSIT, FULL hoặc REMAINING và chưa hoàn trùng", "áp dụng chính sách 100%, 80% hoặc giữ cọc theo thời điểm và bên hủy", "", "Số tiền hoàn bằng 0; hệ thống chỉ hủy booking và không tạo Refund.", extra=["Người thuê hủy trước khi chủ xe duyệt hoặc trước 48 giờ: hoàn 100%.", "Từ 24 đến dưới 48 giờ: hoàn 80%; dưới 24 giờ: giữ cọc.", "Chủ xe hủy: hoàn 100% tiền thuê đã thanh toán hợp lệ."]),
            P("Quy trình tạo hồ sơ hoàn tiền", "Hệ thống", "tạo Refund từ kết quả chính sách hủy", "số tiền hoàn lớn hơn 0 và booking/payment liên quan hợp lệ", "lưu số tiền, phí hủy, chính sách và các payment liên quan", "Refund chuyển sang cần xử lý thủ công (MANUAL_REQUIRED).", "Không đủ điều kiện hoặc đã có hồ sơ tương ứng; hệ thống không tạo trùng."),
            P("Quy trình chủ xe xác nhận đã gửi tiền hoàn", "Chủ xe", "thực hiện hoàn tiền và cung cấp mã tham chiếu, phương thức, ghi chú hoặc minh chứng", "Refund thuộc booking của chủ xe và đang MANUAL_REQUIRED", "ghi nhận thời điểm cùng bằng chứng đã gửi", "Refund chuyển sang đang xử lý (PROCESSING).", "Không đúng quyền hoặc sai trạng thái; hệ thống từ chối."),
            P("Quy trình người thuê xác nhận đã nhận tiền", "Người thuê xe", "kiểm tra khoản hoàn và xác nhận đã nhận đủ", "Refund thuộc booking của mình và đang PROCESSING", "phân bổ số tiền hoàn vào các payment mà không vượt số đã trả", "Refund chuyển sang hoàn thành (SUCCEEDED).", "Người xác nhận không đúng hoặc chưa nhận tiền; Refund chưa hoàn tất."),
        ],
    },
    {
        "title": "Nghiệp vụ bàn giao, trả xe và hoàn tất chuyến thuê",
        "intro": [
            "Giai đoạn thực hiện chuyến thuê bắt đầu khi chủ xe kiểm tra booking, người thuê, giấy tờ và tình trạng thanh toán. Chỉ chủ xe quản lý xe trong booking mới được bàn giao, nhận lại xe, tạo phụ phí và hoàn tất chuyến thuê.",
            "Sau khi đủ điều kiện, xe được bàn giao và booking chuyển sang đang thuê. Khi người thuê trả xe, chủ xe ghi nhận thời gian thực tế và lập biên bản kiểm tra gồm kilomet, nhiên liệu, hình ảnh và tình trạng xe.",
            "Nếu phát sinh trả trễ, thiếu nhiên liệu, vệ sinh hoặc hư hỏng, chủ xe tạo phụ phí có lý do và minh chứng. Booking chỉ hoàn tất sau khi kiểm tra được xóa vướng mắc, không còn payment bắt buộc và không còn phụ phí chờ xử lý.",
        ],
        "processes": [
            P("Quy trình kiểm tra booking trước khi bàn giao", "Chủ xe", "mở booking chuẩn bị bàn giao", "booking thuộc xe mình, đã duyệt, chưa hủy/từ chối, đúng thời gian và thông tin người thuê hợp lệ", "cho phép tiếp tục kiểm tra thanh toán và bàn giao", "", "Một điều kiện không đạt; chủ xe chưa được bàn giao."),
            P("Quy trình xác nhận phần tiền còn lại", "Chủ xe", "đối chiếu totalPrice với các payment PAID", "remainingAmount và khoản tiền thực tế cần thu", "xác nhận tiền mặt hoặc yêu cầu người thuê thanh toán phần còn lại", "Khi remainingAmount bằng 0, booking đủ điều kiện tài chính.", "Còn khoản bắt buộc chưa xử lý; hệ thống không cho bàn giao."),
            P("Quy trình bàn giao xe", "Chủ xe", "xác nhận người thuê đã nhận xe", "booking đã duyệt, giấy tờ đúng, thanh toán đạt yêu cầu và xe chưa bàn giao", "ghi nhận thời điểm bàn giao và trạng thái xe", "Booking chuyển sang đang thuê (IN_PROGRESS); xe chuyển sang đang được thuê (RENTED).", "Không đủ điều kiện; booking và xe giữ nguyên.", "handover"),
            P("Quy trình nhận xe trả", "Chủ xe", "tiếp nhận xe từ người thuê khi kết thúc chuyến", "booking đang ở trạng thái đang thuê (IN_PROGRESS) và thuộc quyền chủ xe", "ghi nhận thời gian trả thực tế và mở hồ sơ kiểm tra", "Booking chuyển sang kiểm tra sau thuê (RETURN_INSPECTION).", "Booking sai trạng thái hoặc không đúng quyền; hệ thống từ chối."),
            P("Quy trình lập biên bản kiểm tra sau thuê", "Chủ xe", "ghi kilomet, nhiên liệu, hình ảnh và tình trạng xe", "dữ liệu kiểm tra đầy đủ và liên kết đúng booking", "lưu biên bản, xác định xe đạt hay có phát sinh", "Không có vấn đề thì kiểm tra chuyển CLEARED; có vấn đề thì chờ phụ phí.", "Thiếu dữ liệu; biên bản chưa được kết thúc.", "return_inspection"),
            P("Quy trình tạo phụ phí phát sinh", "Chủ xe", "chọn loại phí, số tiền, lý do và minh chứng", "booking đang kiểm tra sau thuê và khoản phí hợp lệ", "tạo Extra Charge gắn booking và biên bản", "Phụ phí chuyển sang chờ xử lý (PENDING); booking chuyển sang chờ xử lý phụ phí (AWAITING_EXTRA_CHARGE).", "Sai giai đoạn, số tiền hoặc thiếu lý do; hệ thống không tạo.", "extra_charge"),
            P("Quy trình người thuê thanh toán phụ phí", "Người thuê xe", "chọn phụ phí và phương thức thanh toán", "phụ phí thuộc booking của mình, đang chờ xử lý (PENDING) và số tiền khớp", "tạo payment loại EXTRA_CHARGE và ghi nhận kết quả", "Khi payment đã thanh toán (PAID), phụ phí chuyển sang đã thanh toán (PAID).", "Thanh toán thất bại hoặc phụ phí không hợp lệ; phụ phí vẫn chờ."),
            P("Quy trình chủ xe xác nhận hoặc hủy phụ phí", "Chủ xe", "xác nhận tiền mặt đã thu hoặc hủy khoản phí không còn hợp lệ", "phụ phí thuộc xe mình và đang chờ xử lý (PENDING)", "cập nhật kết quả phụ phí, sau đó kiểm tra các khoản còn lại", "Phụ phí chuyển sang đã thanh toán (PAID) hoặc đã hủy (CANCELLED); hết khoản chờ thì biên bản có thể chuyển sang đã xử lý (CLEARED).", "Phụ phí đã xử lý hoặc không thuộc quyền; hệ thống từ chối."),
            P("Quy trình hoàn tất chuyến thuê", "Chủ xe", "yêu cầu kết thúc booking sau khi nhận xe", "có biên bản đã xử lý (CLEARED), payment đã đủ và không còn phụ phí chờ xử lý (PENDING)", "hoàn tất booking, đồng bộ hợp đồng và giải phóng xe", "Booking chuyển sang hoàn tất (COMPLETED).", "Còn kiểm tra, payment hoặc phụ phí chờ; hệ thống chưa hoàn tất.", "complete"),
        ],
    },
    {
        "title": "Nghiệp vụ hợp đồng, lịch sử và đánh giá",
        "intro": [
            "Hợp đồng là căn cứ gắn trực tiếp với booking và phản ánh người thuê, chủ xe, xe, thời gian, địa điểm, giá cùng nghĩa vụ thanh toán. Code cho phép người thuê tạo hợp đồng ngay sau khi chủ xe duyệt booking và tiếp tục sử dụng hợp đồng trong quá trình thanh toán hoặc chuyến thuê.",
            "Lịch sử booking và payment được hiển thị theo vai trò: người thuê xem dữ liệu do mình tạo, BUSINESS xem dữ liệu từ xe doanh nghiệp và USER ký gửi xem dữ liệu từ xe của mình. Chỉ các bên liên quan được xem hợp đồng.",
            "Sau khi booking hoàn tất, người thuê được tạo tối đa một đánh giá. Chủ xe có thể phản hồi hoặc báo cáo đánh giá thuộc xe mình; ADMIN xem báo cáo và quyết định ẩn hoặc hiển thị lại nội dung.",
        ],
        "processes": [
            P("Quy trình tạo hoặc hoàn thiện hợp đồng", "Người thuê xe", "yêu cầu tạo hợp đồng cho booking đã được duyệt", "booking thuộc người thuê và ở OWNER_APPROVED, PAYMENT_PENDING, PAID, IN_PROGRESS hoặc trạng thái cũ tương ứng", "tạo một hợp đồng duy nhất và đồng bộ giá, cọc, tiền đã trả, còn lại", "Hợp đồng ở trạng thái hoạt động (ACTIVE).", "Booking chưa được duyệt, đã hủy/hoàn tất/no-show hoặc đã có hợp đồng; hệ thống không tạo trùng."),
            P("Quy trình xem hợp đồng", "Người thuê xe hoặc Chủ xe", "chọn hợp đồng liên quan đến booking", "người xem là bên liên quan và hợp đồng tồn tại", "hiển thị nội dung hợp đồng cùng tình trạng thanh toán", "", "Không đúng quyền hoặc không tồn tại; hệ thống không cung cấp dữ liệu."),
            P("Quy trình xem lịch sử booking", "Người dùng", "lọc booking theo mã, xe, trạng thái hoặc thời gian", "vai trò và quyền sở hữu xác định đúng phạm vi", "hiển thị các mốc tạo, duyệt, thanh toán, bàn giao, trả xe và kết thúc", "", "Không có dữ liệu hoặc booking ngoài phạm vi; hệ thống không hiển thị."),
            P("Quy trình xem lịch sử thanh toán", "Người dùng", "xem payment của các booking liên quan", "payment gắn booking thuộc quyền người xem", "hiển thị loại, phương thức, trạng thái, số tiền và thời gian", "", "Payment không thuộc quyền hoặc không tồn tại; hệ thống từ chối."),
            P("Quy trình người thuê đánh giá chuyến thuê", "Người thuê xe", "nhập điểm và nội dung đánh giá", "booking thuộc người thuê, đã COMPLETED và chưa có review", "tạo review gắn booking và xe", "", "Booking chưa hoàn tất hoặc đã có review; hệ thống không tạo.", "review"),
            P("Quy trình chủ xe phản hồi đánh giá", "Chủ xe", "nhập phản hồi cho review", "review thuộc xe mình quản lý và đang được phép hiển thị", "lưu phản hồi mà không thay đổi đánh giá gốc", "", "Review không thuộc quyền; hệ thống từ chối."),
            P("Quy trình chủ xe báo cáo đánh giá vi phạm", "Chủ xe", "gửi lý do báo cáo review", "review thuộc xe mình và lý do báo cáo hợp lệ", "đánh dấu review cần ADMIN xem xét", "", "Không đúng quyền hoặc thiếu lý do; hệ thống không ghi nhận báo cáo."),
            P("Quy trình ADMIN xử lý đánh giá bị báo cáo", "ADMIN", "xem review, lý do và dữ liệu liên quan", "báo cáo tồn tại và quyết định quản trị có căn cứ", "ẩn review vi phạm hoặc hiển thị lại review phù hợp", "", "Chưa đủ căn cứ; ADMIN giữ nguyên trạng thái và tiếp tục xem xét."),
        ],
    },
    {
        "title": "Nghiệp vụ thông báo và theo dõi hệ thống",
        "intro": [
            "Thông báo giúp các bên theo dõi booking, thanh toán, hoàn tiền, bàn giao, trả xe, phụ phí, xe được duyệt và đánh giá. Mỗi thông báo được gắn với người nhận và dữ liệu liên quan nên chỉ người có quyền mới được xem hoặc xử lý.",
            "Trung tâm việc cần làm tập hợp các tác vụ đang chờ theo vai trò, chẳng hạn duyệt xe, duyệt booking, thanh toán, bàn giao, nhận xe trả, hoàn tiền hoặc xử lý đánh giá. Dashboard dùng một quy trình chung nhưng tổng hợp phạm vi dữ liệu khác nhau cho ADMIN, BUSINESS và USER ký gửi.",
        ],
        "processes": [
            P("Quy trình tạo thông báo theo sự kiện nghiệp vụ", "Hệ thống", "nhận biết sự kiện cần thông báo", "người nhận, vai trò, loại sự kiện và dữ liệu liên kết hợp lệ", "tạo thông báo cho đúng người nhận", "", "Không xác định được người nhận hoặc sự kiện không hợp lệ; hệ thống không tạo."),
            P("Quy trình xem thông báo", "Người dùng", "yêu cầu danh sách hoặc chi tiết thông báo", "thông báo thuộc đúng tài khoản và chưa bị xóa mềm", "hiển thị nội dung cùng liên kết nghiệp vụ người dùng được quyền xem", "", "Thông báo không thuộc tài khoản; hệ thống không hiển thị."),
            P("Quy trình đánh dấu thông báo đã đọc", "Người dùng", "đánh dấu một hoặc tất cả thông báo của mình", "phạm vi thông báo thuộc người dùng", "ghi nhận trạng thái đã đọc và thời điểm đọc", "", "Thông báo không thuộc quyền; hệ thống bỏ qua."),
            P("Quy trình xóa thông báo", "Người dùng", "chọn xóa thông báo của mình", "thông báo thuộc tài khoản và chưa bị xóa", "xóa mềm thông báo nhưng không ảnh hưởng dữ liệu nghiệp vụ", "", "Thông báo không thuộc quyền; hệ thống từ chối."),
            P("Quy trình xem trung tâm việc cần làm", "Người dùng", "mở danh sách tác vụ đang chờ", "vai trò và quyền trên từng booking, xe, refund hoặc review", "tổng hợp đúng tác vụ người dùng có thể xử lý", "", "Không có tác vụ phù hợp; hệ thống hiển thị trạng thái rỗng."),
            P("Quy trình xem dashboard theo vai trò", "Người dùng", "mở Dashboard", "vai trò ADMIN, BUSINESS hoặc USER ký gửi và phạm vi dữ liệu", "tổng hợp xe, booking, payment, doanh thu, đánh giá và việc cần làm", "", "Không có dữ liệu hoặc bộ lọc không hợp lệ; hệ thống hiển thị giá trị rỗng."),
        ],
        "dashboard": [
            ("ADMIN", "Toàn hệ thống", "Người dùng, BUSINESS, xe, booking, thanh toán/doanh thu, xe chờ duyệt, đánh giá cần quản trị."),
            ("BUSINESS", "Xe thuộc doanh nghiệp", "Xe, booking mới/đang thuê/hoàn tất, doanh thu từ payment hợp lệ, đánh giá và tác vụ chờ."),
            ("USER ký gửi", "Xe do USER đăng", "Xe đã đăng/chờ duyệt, booking, doanh thu liên quan, đánh giá, bàn giao, trả xe và phụ phí."),
        ],
    },
    {
        "title": "Quy trình thuê xe tổng thể",
        "intro": [
            "Quy trình thuê xe tổng thể liên kết toàn bộ nghiệp vụ từ tìm kiếm đến đánh giá. Người thuê xe chọn địa điểm, thời gian và tiêu chí; hệ thống kiểm tra xe, quyền đặt và lịch thuê trước khi cho giữ xe hoặc tạo booking.",
            "Booking được gửi đến đúng chủ xe để duyệt hoặc từ chối. Sau khi được duyệt, người thuê tạo hoặc hoàn thiện hợp đồng và thực hiện nghĩa vụ thanh toán; chủ xe chỉ bàn giao khi thông tin và payment đáp ứng điều kiện.",
            "Trong chuyến thuê, booking ở trạng thái đang thuê. Khi xe được trả, chủ xe lập biên bản kiểm tra, xử lý phụ phí nếu có và chỉ hoàn tất khi không còn nghĩa vụ chờ. Người thuê sau đó có thể đánh giá chuyến thuê.",
            "Các nhánh ngoại lệ gồm xe hoặc thời gian không hợp lệ, chủ xe từ chối, thanh toán thất bại, booking bị hủy và có thể hoàn tiền, người thuê không đến nhận xe hoặc xe trả có phụ phí.",
        ],
        "processes": [
            P("Quy trình thuê xe tổng thể", "Người thuê xe", "tìm xe, chọn thời gian, xem chi tiết và gửi yêu cầu thuê", "xe hợp lệ, người thuê có quyền đặt, lịch trống và booking được chủ xe duyệt", "hoàn thiện hợp đồng, thanh toán, bàn giao, trả xe, kiểm tra và đánh giá", "Booking lần lượt đi qua REQUESTED, OWNER_APPROVED, PAID, IN_PROGRESS, RETURN_INSPECTION và COMPLETED khi đủ điều kiện.", "Từ chối, thanh toán thất bại, hủy/hoàn tiền, NO_SHOW hoặc phụ phí được chuyển sang nhánh xử lý tương ứng.", "overall", extra=["Chủ xe kiểm tra điều kiện trước khi bàn giao.", "Người thuê trả xe; chủ xe kiểm tra và tạo phụ phí nếu có.", "Chỉ khi mọi nghĩa vụ đã xử lý, booking mới hoàn tất và được đánh giá."]),
        ],
    },
]


DIAGRAM_TITLES = {
    "register": "Đăng ký và xác thực tài khoản",
    "verify_email": "Xác thực email bằng OTP",
    "login": "Đăng nhập",
    "business_account": "ADMIN cấp tài khoản BUSINESS",
    "add_car": "Thêm xe mới",
    "car_approval": "ADMIN duyệt hoặc từ chối xe",
    "location": "Quản lý vị trí nhận xe",
    "search": "Tìm kiếm và chọn xe",
    "cart_hold": "Thêm và giữ xe trong giỏ",
    "create_booking": "Tạo booking",
    "approve_booking": "Duyệt booking",
    "cancel_refund": "Hủy booking và hoàn tiền",
    "no_show": "Xử lý khách không đến nhận xe",
    "payment": "Thanh toán",
    "handover": "Bàn giao xe",
    "return_inspection": "Trả xe và kiểm tra sau thuê",
    "extra_charge": "Tạo và xử lý phụ phí",
    "complete": "Hoàn tất chuyến thuê",
    "review": "Đánh giá chuyến thuê",
    "overall": "Thuê xe tổng thể",
}


def font(path, size):
    return ImageFont.truetype(str(path), size)


REG = Path(r"C:\Windows\Fonts\times.ttf")
BOLD = Path(r"C:\Windows\Fonts\timesbd.ttf")


def wrap(draw, text, fnt, width):
    words = text.split()
    lines, line = [], ""
    for word in words:
        test = f"{line} {word}".strip()
        if draw.textbbox((0, 0), test, font=fnt)[2] <= width:
            line = test
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def centered_text(draw, box, text, fnt, fill="black"):
    x1, y1, x2, y2 = box
    lines = wrap(draw, text, fnt, x2 - x1 - 34)
    line_h = fnt.size + 7
    y = (y1 + y2 - line_h * len(lines)) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=fnt)
        draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, y), line, font=fnt, fill=fill)
        y += line_h


def upper_first(text):
    return text[:1].upper() + text[1:] if text else text


def arrow(draw, a, b, label=None):
    draw.line([a, b], fill="black", width=4)
    angle = math.atan2(b[1] - a[1], b[0] - a[0])
    for delta in (2.65, -2.65):
        p = (b[0] + 16 * math.cos(angle + delta), b[1] + 16 * math.sin(angle + delta))
        draw.line([b, p], fill="black", width=4)
    if label:
        fnt = font(REG, 26)
        mx, my = (a[0] + b[0]) / 2, (a[1] + b[1]) / 2
        label_width = max(88, draw.textbbox((0, 0), label, font=fnt)[2] + 28)
        if abs(b[0] - a[0]) > abs(b[1] - a[1]):
            label_box = (mx - label_width / 2, my - 46, mx + label_width / 2, my - 4)
        else:
            label_box = (mx + 12, my - 22, mx + 12 + label_width, my + 22)
        draw.rectangle(label_box, fill="white")
        centered_text(draw, label_box, label, fnt)


def make_diagram(proc):
    key = proc["diagram"]
    path = DIAGRAM_DIR / f"{key}.png"
    image = Image.new("RGB", (1600, 1120), "white")
    draw = ImageDraw.Draw(image)
    title_font = font(BOLD, 34)
    text_font = font(REG, 28)
    small_font = font(REG, 24)
    centered_text(draw, (80, 30, 1520, 85), DIAGRAM_TITLES[key], title_font)

    main_x1, main_x2 = 360, 1160
    y = 105
    actor_box = (510, y, 1010, y + 80)
    draw.rounded_rectangle(actor_box, radius=38, outline="black", width=4)
    centered_text(draw, actor_box, proc["actor"], text_font)
    y += 115
    action_box = (main_x1, y, main_x2, y + 105)
    draw.rectangle(action_box, outline="black", width=4)
    centered_text(draw, action_box, upper_first(proc["action"]), text_font)
    arrow(draw, (760, actor_box[3]), (760, action_box[1]))

    y += 145
    diamond = [(760, y), (1160, y + 100), (760, y + 200), (360, y + 100)]
    draw.polygon(diamond, outline="black", fill="white")
    centered_text(draw, (425, y + 30, 1095, y + 170), upper_first(proc["check"]), small_font)
    arrow(draw, (760, action_box[3]), (760, y))

    invalid_box = (1210, y + 45, 1550, y + 175)
    draw.rectangle(invalid_box, outline="black", width=4)
    centered_text(draw, invalid_box, "Từ chối xử lý và thông báo lý do", small_font)
    arrow(draw, (1160, y + 100), (1210, y + 100), "Không")

    y += 245
    success_box = (main_x1, y, main_x2, y + 115)
    draw.rectangle(success_box, outline="black", width=4)
    centered_text(draw, success_box, upper_first(proc["success"]), text_font)
    arrow(draw, (760, diamond[2][1]), (760, success_box[1]), "Có")
    y += 155
    result_text = proc["state"] or "Ghi nhận kết quả nghiệp vụ"
    result_box = (main_x1, y, main_x2, y + 110)
    draw.rectangle(result_box, outline="black", width=4)
    centered_text(draw, result_box, result_text, text_font)
    arrow(draw, (760, success_box[3]), (760, result_box[1]))
    y += 150
    end_box = (500, y, 1020, y + 78)
    draw.rounded_rectangle(end_box, radius=38, outline="black", width=4)
    centered_text(draw, end_box, "Thông báo kết quả cho bên liên quan", text_font)
    arrow(draw, (760, result_box[3]), (760, end_box[1]))
    image.save(path, dpi=(180, 180))
    return path


def set_font(run, size=13, bold=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        run._element.rPr.rFonts.set(qn(key), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clear_body(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype in {RT.IMAGE, RT.HYPERLINK}:
            doc.part.drop_rel(rid)


def clear_part(part):
    for paragraph in list(part.paragraphs):
        p = paragraph._element
        for child in list(p):
            p.remove(child)
    for rid, rel in list(part.part.rels.items()):
        if rel.reltype == RT.IMAGE:
            part.part.drop_rel(rid)


def configure_style(style, size, bold=False, before=0, after=3, keep=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = BLACK
    rpr = style._element.get_or_add_rPr()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rpr.rFonts.set(qn(key), FONT)
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = 1.5
    fmt.keep_with_next = keep
    fmt.widow_control = True


def add_update_fields(doc):
    settings = doc.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


def add_field(paragraph, instruction, cached=""):
    run = paragraph.add_run()
    set_font(run)
    for kind, value in (("begin", None), (None, instruction), ("separate", None), (None, cached), ("end", None)):
        if kind:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        else:
            node = OxmlElement("w:instrText" if value == instruction else "w:t")
            node.set(qn("xml:space"), "preserve")
            node.text = value
        run._r.append(node)


def add_page_number(section):
    clear_part(section.footer)
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, " PAGE ", "1")


def prune_unused_header_footer_relationships(doc):
    used = {
        node.get(qn("r:id"))
        for node in doc.sections[0]._sectPr
        if node.tag in {qn("w:headerReference"), qn("w:footerReference")}
    }
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype in {RT.HEADER, RT.FOOTER} and rid not in used:
            doc.part.drop_rel(rid)


def add_heading_numbering(doc):
    numbering = doc.part.numbering_part.element
    aids = [int(e.get(qn("w:abstractNumId"))) for e in numbering.findall(qn("w:abstractNum"))]
    nids = [int(e.get(qn("w:numId"))) for e in numbering.findall(qn("w:num"))]
    aid, nid = max(aids, default=0) + 1, max(nids, default=0) + 1
    abstract = OxmlElement("w:abstractNum"); abstract.set(qn("w:abstractNumId"), str(aid))
    lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), "0")
    for tag, val in (("w:start", "1"), ("w:numFmt", "decimal"), ("w:lvlText", "2.4.1.%1"), ("w:suff", "space")):
        node = OxmlElement(tag); node.set(qn("w:val"), val); lvl.append(node)
    abstract.append(lvl); numbering.append(abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(nid))
    ref = OxmlElement("w:abstractNumId"); ref.set(qn("w:val"), str(aid)); num.append(ref); numbering.append(num)
    return nid


def apply_num(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId"); nid.set(qn("w:val"), str(num_id))
    numpr.extend([ilvl, nid]); ppr.append(numpr)


def add_bullet_numbering(doc):
    numbering = doc.part.numbering_part.element
    aids = [int(e.get(qn("w:abstractNumId"))) for e in numbering.findall(qn("w:abstractNum"))]
    nids = [int(e.get(qn("w:numId"))) for e in numbering.findall(qn("w:num"))]
    aid, nid = max(aids, default=0) + 1, max(nids, default=0) + 1
    abstract = OxmlElement("w:abstractNum"); abstract.set(qn("w:abstractNumId"), str(aid))
    multi = OxmlElement("w:multiLevelType"); multi.set(qn("w:val"), "multilevel"); abstract.append(multi)
    for level, marker, left, hanging in ((0, "●", 720, 360), (1, "○", 1440, 360)):
        lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), str(level))
        for tag, val in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", marker), ("w:suff", "space")):
            node = OxmlElement(tag); node.set(qn("w:val"), val); lvl.append(node)
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind"); ind.set(qn("w:left"), str(left)); ind.set(qn("w:hanging"), str(hanging)); ppr.append(ind)
        lvl.append(ppr)
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts"); fonts.set(qn("w:ascii"), FONT); fonts.set(qn("w:hAnsi"), FONT); rpr.append(fonts)
        lvl.append(rpr); abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(nid))
    ref = OxmlElement("w:abstractNumId"); ref.set(qn("w:val"), str(aid)); num.append(ref); numbering.append(num)
    return nid


def add_step(doc, text, level, bullet_num):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_together = True
    ppr = p._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), str(level))
    nid = OxmlElement("w:numId"); nid.set(qn("w:val"), str(bullet_num))
    numpr.extend([ilvl, nid]); ppr.append(numpr)
    run = p.add_run(text); set_font(run)
    return p


def steps_for(proc):
    result = [
        (0, f'{proc["actor"]} {proc["action"]}.'),
        (0, f'Hệ thống kiểm tra {proc["check"]}.'),
        (1, f'Nếu không hợp lệ: {proc["invalid"]}'),
        (1, "Nếu hợp lệ: hệ thống tiếp tục xử lý quy trình."),
        (0, f'Hệ thống {proc["success"]}.'),
    ]
    if proc["state"]:
        result.append((0, proc["state"]))
    result.extend((0, item) for item in proc["extra"])
    result.append((0, "Hệ thống hiển thị hoặc gửi thông báo kết quả cho các bên liên quan."))
    return result[:10]


def add_caption(doc, title, number):
    p = doc.add_paragraph()
    p.style = doc.styles["Caption"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = False
    r = p.add_run("Hình "); set_font(r, bold=True)
    add_field(p, " SEQ Hình \\* ARABIC ", str(number))
    r = p.add_run(f": Quy trình {title[0].lower() + title[1:]}"); set_font(r, bold=True)


def add_dashboard_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    if "Table Grid" in [style.name for style in doc.styles]:
        table.style = "Table Grid"
    table.autofit = False
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:color"), "000000")
        borders.append(border)
    table._tbl.tblPr.append(borders)
    widths = [Cm(3), Cm(4), Cm(9)]
    headers = ["Vai trò", "Phạm vi dữ liệu", "Nội dung dashboard"]
    for i, cell in enumerate(table.rows[0].cells):
        cell.width = widths[i]; cell.text = headers[i]; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            for r in p.runs: set_font(r, size=12, bold=True)
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].width = widths[i]; cells[i].text = value; cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if i == 2 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.2
                for r in p.runs: set_font(r, size=12)
    doc.add_paragraph()


def build():
    doc = Document(REFERENCE)
    clear_body(doc)
    section = doc.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.left_margin = Cm(3); section.right_margin = Cm(2)
    section.top_margin = Cm(2); section.bottom_margin = Cm(2)
    section.header_distance = Cm(1.2); section.footer_distance = Cm(1.2)
    section.different_first_page_header_footer = False
    clear_part(section.header); add_page_number(section)
    prune_unused_header_footer_relationships(doc)

    configure_style(doc.styles["Normal"], 13, after=3)
    doc.styles["Normal"].paragraph_format.first_line_indent = Cm(1)
    doc.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    configure_style(doc.styles["Heading 1"], 14, bold=True, before=10, after=6, keep=True)
    configure_style(doc.styles["Heading 2"], 13, bold=True, before=8, after=3, keep=True)
    configure_style(doc.styles["Heading 3"], 13, bold=True, before=8, after=3, keep=True)
    if "Caption" not in [style.name for style in doc.styles]:
        doc.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    configure_style(doc.styles["Caption"], 13, bold=True, before=4, after=6, keep=False)
    heading_num = add_heading_numbering(doc)
    bullet_num = add_bullet_numbering(doc)
    add_update_fields(doc)

    spacer = doc.add_paragraph(); spacer.paragraph_format.space_after = Pt(115)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CÁC QUY TRÌNH NGHIỆP VỤ LỚN\nCỦA HỆ THỐNG BQDRIVE"); set_font(r, size=16, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Trình bày theo cấu trúc quy trình và sơ đồ nghiệp vụ"); set_font(r, italic=True)
    doc.add_page_break()

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MỤC LỤC"); set_font(r, size=14, bold=True)
    toc = doc.add_paragraph(); toc.paragraph_format.first_line_indent = Cm(0)
    add_field(toc, ' TOC \\o "1-2" \\h \\z \\u ', "Mục lục được cập nhật khi mở tài liệu.")
    doc.add_page_break()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MỤC LỤC HÌNH"); set_font(r, size=14, bold=True)
    lof = doc.add_paragraph(); lof.paragraph_format.first_line_indent = Cm(0)
    add_field(lof, ' TOC \\h \\z \\c "Hình" ', "Mục lục hình được cập nhật khi mở tài liệu.")
    doc.add_page_break()

    for text in (
        "Các nghiệp vụ nhỏ trong báo cáo BQDrive được tổ chức lại thành tám nhóm lớn theo vòng đời xử lý. Cách trình bày gồm phần giới thiệu bằng văn xuôi, các quy trình cụ thể, bước chính, nhánh điều kiện và sơ đồ ở những luồng quan trọng.",
        "Việc tổ chức lại không thay đổi vai trò, trạng thái, điều kiện tài chính hoặc cấu trúc dữ liệu hiện có. Nội dung nghiệp vụ được đối chiếu với báo cáo gốc và code backend; những điểm chưa thể kết luận được tách riêng ở cuối tài liệu.",
    ):
        p = doc.add_paragraph(); r = p.add_run(text); set_font(r)

    fig_no = 0
    diagram_processes = []
    no_diagram = []
    for gi, group in enumerate(GROUPS):
        if gi > 0:
            doc.add_page_break()
        h = doc.add_paragraph(style="Heading 1"); apply_num(h, heading_num)
        r = h.add_run(group["title"]); set_font(r, size=14, bold=True)
        for intro in group["intro"]:
            p = doc.add_paragraph(); r = p.add_run(intro); set_font(r)
        for proc in group["processes"]:
            h = doc.add_paragraph(proc["title"], style="Heading 2")
            for r in h.runs: set_font(r, bold=True)
            steps = steps_for(proc)
            for level, text in steps:
                add_step(doc, text, level, bullet_num)
            if proc["diagram"]:
                diagram = make_diagram(proc)
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str(diagram), width=Cm(15.2))
                fig_no += 1
                add_caption(doc, DIAGRAM_TITLES[proc["diagram"]], fig_no)
                diagram_processes.append(proc["title"])
            else:
                no_diagram.append(proc["title"])
        if group.get("dashboard"):
            h = doc.add_paragraph("Bảng so sánh dashboard theo vai trò", style="Heading 2")
            for r in h.runs: set_font(r, bold=True)
            add_dashboard_table(doc, group["dashboard"])

    doc.add_page_break()
    h = doc.add_paragraph("Nội dung cần xác nhận", style="Heading 1")
    # Deliberately unnumber this closing note.
    confirmations = [
        "Code dùng khoảng chờ 30 phút sau giờ nhận xe trước khi cho phép xử lý NO_SHOW; cần xác nhận có đưa con số này vào chính sách công bố cho người dùng hay chỉ dùng như quy tắc vận hành.",
        "Một số truy vấn xe công khai chấp nhận cả trạng thái RENTED để hiển thị tình trạng; cần xác nhận giao diện công khai có nên ẩn hoàn toàn xe đang thuê hay vẫn hiển thị nhưng không cho đặt.",
        "Hoàn tiền hiện được theo dõi bằng Refund và xác nhận thủ công giữa chủ xe với người thuê; chưa có luồng tự động hoàn tiền trực tiếp qua MoMo hoặc VNPay.",
        "Việc hiển thị ngày lễ trên lịch và bộ lọc thời gian của dashboard phụ thuộc khả năng hiện tại của giao diện; backend có dữ liệu phục vụ nhưng tài liệu chưa đủ căn cứ xác nhận mọi màn hình đã triển khai.",
    ]
    for item in confirmations:
        add_step(doc, item, 0, bullet_num)

    core = doc.core_properties
    core.title = "Các quy trình nghiệp vụ BQDrive trình bày theo mẫu"
    core.subject = "Tổng hợp tám nhóm nghiệp vụ BQDrive"
    core.author = "BQDrive"
    doc.save(OUTPUT)
    summary = {
        "output": str(OUTPUT),
        "groups": [g["title"] for g in GROUPS],
        "process_count": sum(len(g["processes"]) for g in GROUPS),
        "diagram_count": len(diagram_processes),
        "diagram_processes": diagram_processes,
        "no_diagram_processes": no_diagram,
        "confirmations": confirmations,
    }
    (WORK / "build_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build()
