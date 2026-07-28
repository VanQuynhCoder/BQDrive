from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import re

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r"D:\DoAn_LuanVanTotNghiep")
OUTPUT = ROOT / "BQDrive_Cac_nghiep_vu_trinh_bay_theo_mau.docx"
SOURCES = [
    Path(r"D:\Downloads\BAOCAO_LVTN_demo (4).docx"),
    ROOT / "BQDrive_Cac_nghiep_vu_lon.docx",
    Path(r"C:\Users\bom13\Downloads\VuHoangUng-1.docx"),
]


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def style_name(paragraph):
    return paragraph.style.name if paragraph.style is not None else ""


def has_numbering(paragraph):
    p_pr = paragraph._p.pPr
    return p_pr is not None and p_pr.numPr is not None


def main():
    doc = Document(OUTPUT)
    paragraphs = list(doc.paragraphs)
    headings_1 = [p for p in paragraphs if style_name(p) == "Heading 1"]
    headings_2 = [p for p in paragraphs if style_name(p) == "Heading 2"]
    process_headings = [p for p in headings_2 if p.text.strip().startswith("Quy trình ")]
    numbered_h1 = [p for p in headings_1 if has_numbering(p)]

    step_counts = {}
    for i, paragraph in enumerate(paragraphs):
        if paragraph not in process_headings:
            continue
        count = 0
        for following in paragraphs[i + 1:]:
            if style_name(following) in {"Heading 1", "Heading 2"}:
                break
            if has_numbering(following):
                count += 1
        step_counts[paragraph.text.strip()] = count

    text = "\n".join(p.text for p in paragraphs)
    required_terms = [
        "REQUESTED", "OWNER_APPROVED", "PAYMENT_PENDING", "PAID",
        "IN_PROGRESS", "RETURN_INSPECTION", "AWAITING_EXTRA_CHARGE",
        "COMPLETED", "CANCELLED", "REJECTED", "NO_SHOW",
        "MoMo", "VNPay", "ngày lễ", "bản đồ", "phụ phí", "hoàn tiền",
    ]
    prohibited_terms = ["đơn hàng", "UseCase", "Usecase", "localStorage", "middleware", "router"]

    xml = doc._element.xml
    with ZipFile(OUTPUT) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        settings = archive.read("word/settings.xml").decode("utf-8")
        footer_xml = "\n".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if re.fullmatch(r"word/footer\d+\.xml", name)
        )

    normal = doc.styles["Normal"]
    h1 = doc.styles["Heading 1"]
    h2 = doc.styles["Heading 2"]
    section = doc.sections[0]
    report = {
        "output": str(OUTPUT),
        "output_size": OUTPUT.stat().st_size,
        "source_files": [
            {
                "path": str(path),
                "exists": path.exists(),
                "size": path.stat().st_size if path.exists() else None,
                "modified": path.stat().st_mtime if path.exists() else None,
                "sha256": sha256(path) if path.exists() else None,
            }
            for path in SOURCES
        ],
        "heading1_count": len(headings_1),
        "numbered_group_heading_count": len(numbered_h1),
        "group_headings": [p.text.strip() for p in numbered_h1],
        "heading2_count": len(headings_2),
        "process_heading_count": len(process_headings),
        "unique_process_heading_count": len({p.text.strip() for p in process_headings}),
        "processes_with_fewer_than_4_steps": {
            title: count for title, count in step_counts.items() if count < 4
        },
        "processes_with_more_than_10_steps": {
            title: count for title, count in step_counts.items() if count > 10
        },
        "figure_seq_count": xml.count("SEQ Hình"),
        "toc_field_count": xml.count("TOC "),
        "page_field_count": footer_xml.count("PAGE"),
        "media_count": len(media),
        "update_fields_on_open": '<w:updateFields w:val="true"' in settings,
        "required_terms_missing": [term for term in required_terms if term not in text],
        "prohibited_terms_found": [term for term in prohibited_terms if term in text],
        "page": {
            "width_cm": round(section.page_width.cm, 2),
            "height_cm": round(section.page_height.cm, 2),
            "left_cm": round(section.left_margin.cm, 2),
            "right_cm": round(section.right_margin.cm, 2),
            "top_cm": round(section.top_margin.cm, 2),
            "bottom_cm": round(section.bottom_margin.cm, 2),
        },
        "styles": {
            "normal_font": normal.font.name,
            "normal_size_pt": normal.font.size.pt if normal.font.size else None,
            "normal_spacing": normal.paragraph_format.line_spacing,
            "heading1_font": h1.font.name,
            "heading1_size_pt": h1.font.size.pt if h1.font.size else None,
            "heading2_font": h2.font.name,
            "heading2_size_pt": h2.font.size.pt if h2.font.size else None,
        },
        "table_count": len(doc.tables),
    }
    result_path = ROOT / ".doc_work_bqdrive_sample" / "qa_report.json"
    result_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
